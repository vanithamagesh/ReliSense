"""Paper A v2.3 -> v2.4: equal-coverage comparison (Phase 25, Phase 23 networks only), F1 claim removed,
L10 comparison caveat, limitation sentence. Writes clean v2.4 and a red/blue marked copy."""
import sys, copy, re, difflib
from docx import Document
from docx.shared import RGBColor
src, out_clean, out_marked = sys.argv[1:4]

NEW121 = ("To compare abstention at equal coverage, the LOBO recordings of each network with stored per-recording "
 "predictions were ranked by confidence, and the most confident ones were accepted until the coverage of ReliSense was "
 "reached; WDCNN predictions were not stored and are not included. This ranking uses the test recordings and therefore "
 "favours the networks. At 68.6% coverage, ReliSense reached 97.5% selective accuracy with 2.7% false alarms using its "
 "training-only thresholds, and 96.7% with 4.6% under the same ranking as the networks. The fault-order-guided ResNet "
 "reached 90.8% with 12.1% false alarms, and the impulse-response kernel network and the residual network on raw "
 "vibration 65.4% and 63.9%, with 21.5% and 26.9%. The envelope ResNet raised fewer false alarms (7.5%) but missed "
 "17.9% of the faults, against 0.5% for ReliSense, and reached 75.0%. At 88.7% and 78.6% coverage, ReliSense led the "
 "fault-order-guided ResNet by 6.5 and 6.6 points (89.9% and 94.1% against 83.4% and 87.5%). Part of the lower "
 "false-alarm rate of ReliSense comes from referring more healthy recordings: it accepted 45.6% of them, against "
 "61.5%–88.3% for these networks.")

EDITS = {  # paragraph index in v2.3 -> list of (old, new) replacements inside runs
 3: [("version 2.3, 6 October 2026", "version 2.4, 8 October 2026")],
 107: [("Recording accuracy is reported with macro F1-score and the macro area",
        "Recording accuracy is reported, with the macro area")],
 151: [("(LOBO) in this study.", "(LOBO) in this study. The correspondence of the ten L10 combinations of this study "
        "to the published L10 evaluation has not been established, so its 98.3% is listed for reference and not as a "
        "matched comparison.")],
 165: [("code and were not tuned. ", "code and were not tuned. The equal-coverage comparison ranked the test recordings "
        "of the networks and is therefore descriptive. ")],
}

def build(marked):
    d = Document(src); P = d.paragraphs
    old = {i: P[i].text for i in list(EDITS) + [121]}
    for i, reps in EDITS.items():
        for o, n in reps:
            hit = [r for r in P[i].runs if o in r.text]
            assert len(hit) == 1, (i, o); hit[0].text = hit[0].text.replace(o, n)
    # new paragraph after 121
    newp = copy.deepcopy(P[121]._p); P[121]._p.addnext(newp)
    from docx.text.paragraph import Paragraph
    np_ = Paragraph(newp, P[121]._parent)
    for r in np_.runs[1:]: r._r.getparent().remove(r._r)
    np_.runs[0].text = NEW121
    if marked:
        for i in EDITS:
            mark(P[i], old[i])
        for r in np_.runs: r.font.color.rgb = RGBColor(0, 0, 255)
        key = P[0].insert_paragraph_before("")
        a = key.add_run("Revision key (v2.3 → v2.4):  "); 
        b = key.add_run("red strikethrough = deletion"); b.font.color.rgb = RGBColor(255, 0, 0); b.font.strike = True
        key.add_run("   |   ")
        c = key.add_run("blue = addition"); c.font.color.rgb = RGBColor(0, 0, 255)
    return d

def mark(p, oldtext):
    new = p.text; style_run = p.runs[0]
    A = re.findall(r'\S+\s*', oldtext); B = re.findall(r'\S+\s*', new)
    tmpl = copy.deepcopy(style_run._r)
    for r in list(p.runs): r._r.getparent().remove(r._r)
    def add(t, col=None, strike=False):
        r = p.add_run(t); rr = r._r; rpr = tmpl.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr')
        if rpr is not None: rr.insert(0, copy.deepcopy(rpr))
        if col: r.font.color.rgb = col
        if strike: r.font.strike = True
    for op, a1, a2, b1, b2 in difflib.SequenceMatcher(None, A, B, autojunk=False).get_opcodes():
        if op == 'equal': add(''.join(A[a1:a2]))
        if op in ('delete', 'replace'): add(''.join(A[a1:a2]), RGBColor(255, 0, 0), True)
        if op in ('insert', 'replace'): add(''.join(B[b1:b2]), RGBColor(0, 0, 255))

build(False).save(out_clean); build(True).save(out_marked)
print('ok')
