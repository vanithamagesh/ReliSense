"""Paper A v2.5 from the user's edited v2.3 (e3b64cc6 upload): v2.4 changes + phase 26 / 26b results.
python3 make_A25.py <in.docx> <out_clean.docx> <out_marked.docx>
Real numbers only: phase 25 (matched coverage), phase 26 (balanced decision, other fusions, external abstention),
phase 26b (per-bearing tests, operating points). Marked copy: red strikethrough = deletion, blue = addition."""
import sys, copy, re, difflib
from docx import Document
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph
from docx.table import Table

src, out_clean, out_marked = sys.argv[1:4]
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
RED, BLUE = RGBColor(255, 0, 0), RGBColor(0, 0, 255)

# ------------------------------------------------------------------ text edits (index in the input file: [(old, new)])
EDITS = {
 3: [('Paper A, version 2.3, 6 October 2026', 'Paper A, version 2.5, 8 October 2026')],
 4: [('to 2.7%. With', 'to 2.7%. Correcting the decision to equal class priors, without retraining, halved the false alarms '
     'before abstention to 15.6%, at the cost of more missed faults (11.2% instead of 6.6%). With')],
 107: [('Recording accuracy is reported with macro F1-score and the macro area',
        'Recording accuracy is reported, with balanced accuracy, macro F1-score and the macro area')],
 110: [('exceeds every result of this study.', 'exceeds every result of this study. Vieira et al. [15], who followed '
        'bearing-wise splits of the same kind, could not reproduce this value.')],
 111: [(' 95.7%). Against WDCNN', ' 95.7%; balanced accuracy 81.2%; macro F1-score 81.5%). Against WDCNN')],
 112: [('(18.1% against 32.1%), but on L8 it diagnosed', '(18.1% against 32.1%), but it missed 28.0% of the faults, '
        'against 6.6%, and on L8 it diagnosed')],
 133: [('so such a dataset cannot measure transfer to new bearings.', 'so such a dataset cannot measure transfer to new '
        'bearings. Features that identify the individual bearing, rather than its fault, can therefore classify every segment.')],
 143: [('With abstention at κ = 0.7, it also gave the highest selective accuracy (97.5%) with 2.7% false alarms.',
        'With abstention at κ = 0.7, it gave the highest selective accuracy (97.5%) with 2.7% false alarms, equal to the '
        'combination of the core with the envelope branch. The modulation features thus improved the diagnosis before '
        'abstention but not the accepted diagnoses at this coverage.')],
 151: [('(LOBO) in this study.', '(LOBO) in this study. The correspondence of the ten L10 combinations of this study to the '
        'published L10 evaluation has not been established, so its 98.3% is listed for reference and not as a matched comparison.')],
 165: [('code and were not tuned. ', 'code and were not tuned. The equal-coverage comparison ranked the test recordings of '
        'the networks and is therefore descriptive. '),
       ('Abstention was evaluated on Paderborn only and not on the external datasets. Future work should accept a diagnosis '
        'only when several harmonics of the candidate fault agree. ',
        'On the further datasets, the abstention thresholds needed at least two training bearings per class. Features that '
        'require the agreement of several harmonics did not improve the diagnosis, and the equal-prior decision was evaluated '
        'after the main results were known. '),
       ('collection [46]', 'collection [47]'), ('shift [47], [48]', 'shift [48], [49]'), ('prediction [49]', 'prediction [50]')],
 167: [('at the cost of referring 31% of the recordings.', 'at the cost of referring 31% of the recordings. A decision with '
        'equal class priors, applied without retraining, halved the false alarms before abstention, at the cost of more missed faults.')],
 218: [('[46]', '[47]')], 219: [('[47]', '[48]')], 220: [('[48]', '[49]')], 221: [('[49]', '[50]')],
}

NEW_AFTER = {   # new paragraphs inserted after paragraph i (template: paragraph i)
 121: ["To compare abstention at equal coverage, the LOBO recordings of each network with stored per-recording predictions "
       "were ranked by confidence, and the most confident ones were accepted until the coverage of ReliSense was reached; WDCNN "
       "predictions were not stored and are not included. This ranking uses the test recordings and therefore favours the "
       "networks. At 68.6% coverage, ReliSense reached 97.5% selective accuracy with 2.7% false alarms using its training-only "
       "thresholds, and 96.7% with 4.6% under the same ranking as the networks. The fault-order-guided ResNet reached 90.8% with "
       "12.1% false alarms, and the impulse-response kernel network and the residual network on raw vibration 65.4% and 63.9%, "
       "with 21.5% and 26.9%. The envelope ResNet raised fewer false alarms (7.5%) but missed 17.9% of the faults, against 0.5% "
       "for ReliSense, and reached 75.0%. At 88.7% and 78.6% coverage, ReliSense led the fault-order-guided ResNet by 6.5 and "
       "6.6 points (89.9% and 94.1% against 83.4% and 87.5%). Part of the lower false-alarm rate of ReliSense comes from "
       "referring more healthy recordings: it accepted 45.6% of them, against 61.5%–88.3% for these networks."],
 154: ["Table VIII compares the two branches and ReliSense over the eleven bearing-wise settings of this study. The better "
       "branch differed between datasets: the kinematic branch was better on L8, L10, A2R, CWRU and HUST, the envelope branch "
       "under LOBO and on Ottawa. Each branch alone therefore failed somewhere. The kinematic branch fell 6.3 points below the "
       "envelope branch on Ottawa with one training profile, and the envelope branch fell 11.4 points below the kinematic "
       "branch on L8 at condition 0 and on HUST with one training type. ReliSense was at least as accurate as the better branch "
       "in 9 of the 11 settings and at most 1.6 points below it in the other two. The benefit of averaging is therefore "
       "robustness across datasets rather than a significant gain on any one of them. Three more selective combinations were "
       "tested afterwards on the same splits and did not improve on equal weights. A weight chosen on the inner held-out "
       "bearings gave 83.8% under LOBO and 95.0% on Ottawa with one training profile, against 84.7% and 95.9% for equal "
       "weights, both with the equal-prior decision introduced below. A stacked logistic regression on the branch posteriors "
       "fell to 63.3% on L10. Six harmonic-agreement features, the median and the minimum ratio over the three harmonics of "
       "BPFO, BPFI and BSF, changed the accuracy of the kinematic branch by between −1.9 and +2.3 points."],
 157: ["The balance can also be shifted before abstention, without retraining. Both branches are fitted on training sets in "
       "which the healthy class is the smallest, 6 of 29 bearings under LOBO, so their posteriors carry the class frequencies "
       "of the training bearings. Dividing each posterior by these frequencies and renormalizing gives the decision for equal "
       "class priors [46]. Table IX compares the two decision rules. Under LOBO, equal priors raised the accuracy from 83.7% to "
       "84.7% and the healthy recall from 67.9% to 84.4%, and lowered the false alarms from 32.1% to 15.6%, below the 18.1% of "
       "the envelope ResNet. The missed faults rose from 6.6% to 11.2%, against 28.0% for that network. K004 and K005 were then "
       "diagnosed correctly in 61.3% and 72.5% of their recordings. Dividing by the class frequencies raised to the power "
       "0.25, 0.5 and 0.75 gave intermediate operating points, with false alarms of 27.5%, 22.7% and 19.0% and missed faults of "
       "7.7%, 9.1% and 10.1%. The correction moves errors between classes rather than removing them. The weak real outer-race "
       "damage lost accuracy (KA15 from 68.8% to 52.5%, KA22 from 75.0% to 62.5% and KA30 from 42.5% to 27.5%), and per "
       "bearing the two rules did not differ significantly (6 bearings better, 10 worse; Wilcoxon p = 0.59). On CWRU under "
       "LOSO, equal priors removed all false alarms and raised the accuracy to 86.3%, above both branches; with one training "
       "size, the accuracy fell from 76.0% to 73.3%. All other results in this article use the training priors, which gave "
       "the fewest missed faults."],
 163: ["Abstention was also applied to the further datasets, with thresholds from the same inner loop over the training units "
       "(Table X). Where the inner loop could hold out the same kind of unit as the test, the coverage stayed near its target "
       "and the accepted segments were more accurate. Under HUST LOTO, 71.1% of the segments were accepted with 92.2% "
       "selective accuracy, against 85.0% without abstention; on CWRU under LOSO, 83.2% were accepted with 91.8%, and on "
       "Ottawa 76.0% and 79.1% with 100% and 99.7%. Where the training set contained a single fault size or bearing type, the "
       "inner loop could hold out only loads. Confidences on held-out loads of known bearings exceeded those on new bearings, "
       "so the thresholds were too strict and accepted only 18.5% of the CWRU and 27.2% of the HUST segments. The calibration "
       "therefore needs at least two training bearings per class, as Paderborn provides."],
}

TABLES = {   # inserted after the new paragraph that follows paragraph i: (number, title, header, rows, footnote)
 154: ('TABLE VIII', 'Accuracy of the Two Branches and of ReliSense in All Bearing-Wise Settings (%)',
       ['Setting', 'Kinematic branch', 'Envelope branch', 'ReliSense', 'ReliSense, equal priors'],
       [['Paderborn L8, cond. 0', '95.5', '84.1', '95.9', '95.5'], ['Paderborn L8, all', '87.0', '81.7', '87.0', '88.6'],
        ['Paderborn L10', '77.2', '72.6', '78.7', '78.7'], ['Paderborn A2R', '76.5', '76.3', '77.7', '77.9'],
        ['Paderborn LOBO', '81.6', '82.5', '83.7', '84.7'], ['CWRU LOSO', '84.5', '75.2', '83.2', '86.3'],
        ['CWRU 1SIZE', '76.0', '67.8', '76.0', '73.3'], ['HUST LOTO', '83.3', '81.1', '85.0', '85.0'],
        ['HUST 1TYPE', '73.8', '62.4', '73.9', '73.9'], ['Ottawa LOPO', '95.9', '98.3', '98.3', '98.3'],
        ['Ottawa 1PROF', '91.2', '97.5', '95.9', '95.9'],
        ['At least as accurate as the better branch', '8 of 11', '3 of 11', '9 of 11', '9 of 11'],
        ['Largest shortfall (points)', '6.3', '11.4', '1.6', '2.7']],
       'Paderborn: recordings; further datasets: segments pooled over folds; L10: mean over 10 splits. Equal priors: both '
       'branch posteriors divided by the training class frequencies before averaging (Section VI-C, part 2).'),
 157: ('TABLE IX', 'Decision With Training Priors and With Equal Priors: Accuracy / False Alarms / Missed Faults (%)',
       ['Setting', 'Training priors', 'Equal priors'],
       [['Paderborn L8, cond. 0', '95.9 / 0.0 / 0.0', '95.5 / 0.0 / 1.0'], ['Paderborn L8, all', '87.0 / 43.8 / 0.9', '88.6 / 18.8 / 3.5'],
        ['Paderborn A2R', '77.7 / 65.8 / 0.6', '77.9 / 59.2 / 1.4'], ['Paderborn LOBO', '83.7 / 32.1 / 6.6', '84.7 / 15.6 / 11.2'],
        ['Paderborn LOBO, κ = 0.7', '97.5 / 2.7 / 0.5', '96.4 / 1.7 / 1.9'], ['CWRU LOSO', '83.2 / 16.7 / 16.8', '86.3 / 0.0 / 16.8'],
        ['CWRU 1SIZE', '76.0 / 10.0 / 16.8', '73.3 / 0.0 / 21.0']],
       'LOBO, κ = 0.7: selective accuracy / false alarms / missed faults with training-only thresholds; coverage 68.6% and '
       '69.9%, of the healthy recordings 45.6% and 58.5%. Balanced accuracy under LOBO: 81.2% and 84.7%. L10, HUST and Ottawa '
       'have equal class frequencies in training, so both rules give the same results there.'),
 163: ('TABLE X', 'Abstention on the Further Datasets With Training-Only Thresholds (κ = 0.7, %)',
       ['Protocol', 'Inner held-out unit', 'Accuracy, all segments', 'Coverage', 'Selective accuracy', 'False alarms'],
       [['CWRU LOSO', 'Fault size (healthy: load)', '83.2', '83.2', '91.8', '3.3'], ['CWRU 1SIZE', 'Load', '76.0', '18.5', '83.3', '0.0'],
        ['HUST LOTO', 'Bearing type', '85.0', '71.1', '92.2', '8.3'], ['HUST 1TYPE', 'Load', '73.9', '27.2', '88.8', '0.4'],
        ['Ottawa LOPO', 'Speed profile', '98.3', '76.0', '100.0', '0.0'], ['Ottawa 1PROF', 'Recording', '95.9', '79.1', '99.7', '0.0']],
       'Training priors. False alarms: healthy segments accepted and diagnosed as faulty, as a share of all healthy segments. '
       'If a class has fewer than two training units of the held-out kind, the inner loop holds out loads (CWRU, HUST) or '
       'recordings (Ottawa).'),
}

NEW_REF_AFTER = 217   # after [45]
NEW_REF = ('[46]\tM. Saerens, P. Latinne, and C. Decaestecker, "Adjusting the outputs of a classifier to new a priori '
           'probabilities: A simple procedure," Neural Comput., vol. 14, no. 1, pp. 21–41, 2002, doi: 10.1162/089976602753284446.')
CUT_FROM = 222        # the pasted working note and its table after the reference list are removed


def replace_across_runs(p, old, new):
    runs = p.runs; full = ''.join(r.text for r in runs); k = full.find(old)
    assert k >= 0 and full.count(old) == 1, (old[:60], full[:80])
    pos = 0; first = True; end = k + len(old)
    for r in runs:
        a, b = pos, pos + len(r.text); pos = b
        if b <= k or a >= end: continue
        s, e = max(k, a) - a, min(end, b) - a
        if first: r.text = r.text[:s] + new + r.text[e:]; first = False
        else: r.text = r.text[:s] + r.text[e:]


def new_par_after(anchor_el, template_p, text):
    el = copy.deepcopy(template_p._p)
    rs = el.findall(W + 'r'); keep = max(rs, key=lambda r: len(''.join(t.text or '' for t in r.iter(W + 't'))))
    for r in rs:
        if r is not keep: el.remove(r)
    for tag in ('hyperlink', 'bookmarkStart', 'bookmarkEnd', 'proofErr'):
        for x in el.findall(W + tag): el.remove(x)
    anchor_el.addnext(el); p = Paragraph(el, template_p._parent); p.runs[0].text = text
    return p


def build_table(doc, tmpl_tbl, header, rows):
    t = copy.deepcopy(tmpl_tbl._tbl); n = len(header)
    grid = t.find(W + 'tblGrid'); cols = grid.findall(W + 'gridCol')
    total = sum(int(c.get(W + 'w')) for c in cols); w = total // n
    for c in cols[n:]: grid.remove(c)
    while len(grid.findall(W + 'gridCol')) < n: grid.append(copy.deepcopy(cols[0]))
    for c in grid.findall(W + 'gridCol'): c.set(W + 'w', str(w))
    trs = t.findall(W + 'tr'); hdr, body = trs[0], trs[1]
    for tr in trs[2:]: t.remove(tr)
    for tr in (hdr, body):
        tcs = tr.findall(W + 'tc')
        for tc in tcs[n:]: tr.remove(tc)
        while len(tr.findall(W + 'tc')) < n: tr.append(copy.deepcopy(tr.findall(W + 'tc')[-1]))
        for tc in tr.findall(W + 'tc'):
            tcw = tc.find(W + 'tcPr/' + W + 'tcW')
            if tcw is not None: tcw.set(W + 'w', str(w)); tcw.set(W + 'type', 'dxa')
    t.remove(body)
    for row in rows:
        tr = copy.deepcopy(body); t.append(tr)
    table = Table(t, tmpl_tbl._parent)
    def setc(cell, text):
        ps = cell.paragraphs
        for extra in ps[1:]: extra._p.getparent().remove(extra._p)
        p = ps[0]; rs = p.runs
        if not rs: p.add_run(text); return
        rs[0].text = text
        for r in rs[1:]: r._r.getparent().remove(r._r)
    for j, h in enumerate(header): setc(table.rows[0].cells[j], h)
    for i, row in enumerate(rows):
        for j, v in enumerate(row): setc(table.rows[i + 1].cells[j], v)
    return table


def mark(p, oldtext):
    new = p.text; A = re.findall(r'\S+\s*', oldtext); B = re.findall(r'\S+\s*', new)
    tmpl = copy.deepcopy(max(p.runs, key=lambda r: len(r.text))._r)
    for r in list(p.runs): r._r.getparent().remove(r._r)
    def add(t, col=None, strike=False):
        r = p.add_run(t); rpr = tmpl.find(W + 'rPr')
        if rpr is not None: r._r.insert(0, copy.deepcopy(rpr))
        if col: r.font.color.rgb = col
        if strike: r.font.strike = True
    for op, a1, a2, b1, b2 in difflib.SequenceMatcher(None, A, B, autojunk=False).get_opcodes():
        if op == 'equal': add(''.join(A[a1:a2]))
        if op in ('delete', 'replace'): add(''.join(A[a1:a2]), RED, True)
        if op in ('insert', 'replace'): add(''.join(B[b1:b2]), BLUE)


def blue(p):
    for r in p.runs: r.font.color.rgb = BLUE


def build(marked):
    d = Document(src); P = list(d.paragraphs); body = d.element.body
    # remove the pasted working note (paragraphs from CUT_FROM and the table after the reference list)
    last_ref = P[CUT_FROM - 1]._p; el = last_ref.getnext()
    while el is not None and not el.tag.endswith('sectPr'):
        nxt = el.getnext(); body.remove(el); el = nxt
    old = {i: P[i].text for i in EDITS}
    for i, reps in EDITS.items():
        for o, n in reps: replace_across_runs(P[i], o, n)
    tmpl_tbl = d.tables[5]                                                  # Table VI as format template
    cap_no, cap_title = P[144], P[145]; foot = P[146]
    for i in sorted(NEW_AFTER, reverse=True):
        anchor = P[i]._p; newp = []
        for txt in NEW_AFTER[i]:
            q = new_par_after(anchor, P[i], txt); newp.append(q); anchor = q._p
        if i in TABLES:
            no, title, hdr, rows, fn = TABLES[i]
            q1 = new_par_after(anchor, cap_no, no); q2 = new_par_after(q1._p, cap_title, title)
            tb = build_table(d, tmpl_tbl, hdr, rows); q2._p.addnext(tb._tbl)
            q3 = new_par_after(tb._tbl, foot, fn); newp += [q1, q2, q3]
            if marked:
                for row in tb.rows:
                    for c in row.cells:
                        for pp in c.paragraphs: blue(pp)
        if marked:
            for q in newp: blue(q)
    r = new_par_after(P[NEW_REF_AFTER]._p, P[NEW_REF_AFTER], NEW_REF)
    if marked:
        blue(r)
        for i in EDITS: mark(P[i], old[i])
        key = P[0].insert_paragraph_before('')
        key.add_run('Revision key (v2.3 with author edits → v2.5):  ')
        x = key.add_run('red strikethrough = deletion'); x.font.color.rgb = RED; x.font.strike = True
        key.add_run('   |   '); x = key.add_run('blue = addition'); x.font.color.rgb = BLUE
        key.add_run('   |   The working note after the reference list was removed.')
    return d


build(False).save(out_clean); build(True).save(out_marked); print('ok')
