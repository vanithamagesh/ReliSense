"""Paper A v2.9 from v2.8: shortened to below 10,000 words as counted by Word (all text, tables and references).
Text: paragraphs condensed further; every number in a rewritten paragraph must already occur in the paragraph it
replaces (derived values listed explicitly). Tables: same values, compact notation a/b/c instead of a / b / c; Table II
settings shortened; Table XII without the single-recording column, which repeats Table X. No result, table, figure or
reference removed.  python3 make_A29.py <v2_8.docx> <out_clean.docx> <out_marked.docx>
Marked copy: red strikethrough = deletion, blue = addition (table notation changes are not marked)."""
import sys, copy, re, difflib
from docx import Document
from docx.shared import RGBColor

src, out_clean, out_marked = sys.argv[1:4]
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
RED, BLUE = RGBColor(255, 0, 0), RGBColor(0, 0, 255)

REWRITE = [   # (start of the paragraph, new text [without an italic run-in heading], derived numbers allowed)
 ('Paper A, version 2.8', 'Paper A, version 2.9, 8 October 2026', {'2.9'}),
 ('Abstract—',
  'Abstract—A bearing fault classifier in service must diagnose bearings that it has never seen, although public datasets '
  'contain few physical bearings per class. ReliSense, the physics-informed framework proposed here, reads the squared '
  'envelope spectrum only at the fault orders fixed by bearing geometry and measured shaft speed. A 96-parameter logistic '
  'regression and a random forest on four demodulation bands classify these features, their posteriors are averaged, and '
  'low-confidence recordings are withheld with thresholds set on training bearings only. On 29 Paderborn bearings, which '
  'also served for method development, ReliSense reached 95.9% when trained on artificial and tested on real damage, '
  'against 75.0% for the published ensemble and at most 70.8% for the compared networks. In leave-one-bearing-out '
  'evaluation, it reached 83.7%; abstention raised this to 97.5% at 69% coverage and lowered the false alarms from 32.1% '
  'to 2.7%. Equal class priors lowered the false alarms before abstention to 15.6%, and a decision from ten consecutive '
  'recordings to 6.2%, with 11.2% and 8.7% missed faults. With settings fixed in advance, ReliSense reached 85.0% on '
  'unseen bearing types of the HUST dataset and 83.2% on unseen fault sizes of the CWRU dataset. Healthy bearings that '
  'resemble damage remain the main source of error.', set()),
 # ---------------------------------------------------------------- I
 ('Rolling-element bearings support',
  'Rolling-element bearings support almost every rotating shaft in industrial machinery, and their failure is a leading '
  'cause of unplanned downtime [1]. A localized defect produces a short impact each time a rolling element passes over it, '
  'at a rate fixed by the shaft speed and the bearing geometry, and the envelope spectrum of the excited resonance shows '
  'this rate as distinct lines [2]. Envelope analysis is therefore the industrial reference for bearing diagnosis.', set()),
 ('Data-driven methods now dominate',
  'Data-driven methods now dominate research on this task [3]–[5]. Networks such as WDCNN [6], transformers [7] and '
  'transfer learning [8] learn features directly from vibration, and reviews report accuracies close to 100% [9]. These '
  'accuracies depend on how the data are split: on the Paderborn dataset [10], [11], the splitting method changed accuracy '
  'by more than 40%, and ten of 55 studies were judged likely to be affected by data leakage [12]. When training and test '
  'bearings are kept apart, the published accuracy for transfer from artificial to real damage falls to 62.3%–65.9% for '
  'single classifiers and 75.0% for an ensemble. A monitoring system in service always meets unseen bearings, so this is '
  'the setting that matters.', set()),
 ('Diagnosing an unseen bearing poses',
  'Diagnosing an unseen bearing is difficult for three reasons. Public datasets contain many recordings of few physical '
  'bearings, six or seven per class on Paderborn, so a model can separate its training bearings by properties unrelated '
  'to the damage [13], [14], and the number of training bearings largely decides generalization [15]. Training data may '
  'contain only artificial damage, whereas damage in service grows by fatigue or plastic deformation. Small or '
  'distributed damage produces weak impacts, which lower speed or load weakens further [16], [17].', set()),
 ('Signal-processing methods select',
  'Signal-processing methods select a demodulation band with the kurtogram [18], combined indicators [19], [20] or '
  'cepstral pre-whitening [21], [22], but the selected band is not always the one excited by the defect, and the decision '
  'requires an expert. Data-driven methods learn the whole chain [23]–[26], and transfer from artificial to natural damage '
  'has used data of the target domain [27]. Physics-informed networks add characteristic-frequency priors [28], modal '
  'properties [29], fault-frequency-guided decomposition [30] or interpretable filtering layers [31], [32], and fault '
  'orders have weighted envelope features [33]. Uncertainty estimates and selective prediction aim at trustworthy '
  'diagnoses [34]–[37].', set()),
 ('Despite this progress, two research gaps',
  'Two research gaps remain. Physics-informed models are rarely evaluated on bearings absent from training or compared '
  'with simpler models that use the same physics, so it is not known whether their prior carries over to a new bearing. '
  'Confidence-based rejection is usually calibrated on data of the test domain [38], which is not available for an unseen '
  'bearing.', set()),
 ('This study proposes ReliSense',
  'This study proposes ReliSense, which reads the squared envelope spectrum only at the fault orders fixed by bearing '
  'geometry and measured shaft speed, classifies these features in two branches, and withholds diagnoses below a '
  'threshold set on training bearings only. Physical knowledge thus enters through the input representation rather than '
  'the loss [39]. ReliSense is evaluated under bearing-wise protocols on Paderborn and, with settings fixed in advance, on '
  'three further test rigs.', set()),
 ('Fault-order analysis itself is established',
  'Fault-order analysis is established; the novelty lies in combining fault-order features, two demodulation branches and '
  'an abstention threshold that needs no data of the new bearing. The contributions are as follows.', set()),
 ('A fault-order representation for unseen bearings',
  'A fault-order representation for unseen bearings: every input is read at an order computed from the geometry and the '
  'measured speed, including 16 sideband and higher-harmonic features predicted by the defect kinematics.', set()),
 ('A two-branch architecture',
  'A two-branch architecture: a 96-parameter logistic regression with weights traceable to fault orders, and a random '
  'forest that reads the same peaks in four demodulation bands; averaging lowers the confidence when the two disagree.', set()),
 ('A bearing-wise abstention rule',
  'A bearing-wise abstention rule: the threshold of each test bearing is set for a target coverage by an inner '
  'leave-one-bearing-out loop over the training bearings, without data of the new bearing.', set()),
 # ---------------------------------------------------------------- II
 ('Fig. 1 shows the framework',
  'Fig. 1 shows the framework: Module I converts one 4-s recording into kinematic and envelope features, and Module II classifies them in two branches. The kinematic core was fixed in a preliminary study on 15 of the bearings (K001–K005, KA01, KA04, KA05, KA07, KA15, KI01, KI04, KI05, KI07 and KI14); the modulation features and the envelope branch were added after its Paderborn results were known, and the final configuration was fixed before the further datasets were analyzed.', set()),
 ('Fig. 1. Overall architecture',
  'Fig. 1. Architecture of ReliSense, with the real Module I output for recording N15_M07_F10_KA16_1 and the pooled LOBO '
  'posteriors and confidences.', set()),
 ('Fig. 2. Kinematics of the 6203',
  'Fig. 2. Kinematics of the 6203 bearing. (a) Rolling element between the rings. (b) Load-zone function q(θ) of (4) for '
  'ε = 0.5 and 0.3. (c) Search windows on the order axis; dotted: overlapping pairs.', set()),
 ('Fig. 3. Local defects',
  'Fig. 3. Local defects of a ball bearing. (a)–(c) Rendered outer-raceway, inner-raceway and rolling-element defects '
  '(not to scale). (d)–(f) Simulated impact trains of (3) at 1500 rpm.', set()),
 ('Fig. 5 tests the normalization',
  'Fig. 5 tests the normalization on real recordings at 900 and 1500 rpm. On the order axis, the fault lines stay at the '
  'fault orders. Read at the frequencies of 1500 rpm instead of the measured speed, the first-harmonic feature at 900 rpm '
  'fell from 3.30 to 1.90 for KA16 and from 2.66 to 0.65 for KI18, the level of the healthy K001 (0.48–0.63).', set()),
 ('Fig. 4. Real output of Module I',
  'Fig. 4. Real output of Module I for recording N15_M07_F10_KA16_1 (outer-race damage, 1500 rpm): (1) raw vibration; '
  '(2) amplitude spectrum and band; (3) squared envelope; (4) envelope spectrum with BPFO harmonics; (5) shaft speed; (6) '
  'order spectrum with the healthy K001; (7) peak and background windows; (8) the 15 features.', set()),
 ('Fig. 5. Speed normalization',
  'Fig. 5. Speed normalization on Paderborn recordings at 900 and 1500 rpm. Top: envelope spectra with the first three '
  'fault harmonics at each speed. Middle: the same on the order axis. Bottom: first-harmonic feature (mean ± SD of 20 '
  'recordings per condition).', set()),
 ('The branch has 3 × 31 weights',
  'The branch has 3 × 31 weights and 3 biases, 96 parameters; a linear model was chosen because, with six or seven '
  'bearings per class, every additional degree of freedom can learn the identity of the training bearings. The envelope '
  'branch is a random forest of 500 trees (seed 0) on e. The two posteriors are averaged,', set()),
 ('Fig. 7. LOBO posteriors',
  'Fig. 7. LOBO posteriors of the 2319 recordings in the probability simplex. (a) True classes (grey: H, blue: OR, orange: '
  'IR). (b) Decision at the mean threshold for κ = 0.7; shaded: withheld region.', set()),
 ('Fig. 9 summarizes the procedure',
  'Fig. 9 summarizes the procedure: everything is computed offline from the training bearings, and a new bearing needs '
  'one recording, its shaft speed and its geometry.', set()),
 # ---------------------------------------------------------------- III
 ('The Paderborn dataset was recorded',
  'The Paderborn dataset was recorded on a modular rig in which a motor drives the test-bearing module against a load '
  'motor (Fig. 10), with the housing vibration sampled at 64 kHz. Each bearing was run under four conditions with 20 '
  'recordings of 4 s: N15_M07_F10 (condition 0: 1500 rpm, 0.7 Nm, 1000 N), N09_M07_F10 (900 rpm), N15_M01_F10 (0.1 Nm) '
  'and N15_M07_F04 (400 N). It is the only public dataset with artificial and real damage on many bearings of one type.', set()),
 ('Fig. 10. Paderborn test rig',
  'Fig. 10. Paderborn rig (a) and 6203 test bearing (b), rendered after the dataset description (n = 8, d = 6.75 mm, '
  'D = 28.55 mm).', set()),
 ('Table I lists the bearings',
  'Table I lists the bearings: six healthy, twelve with artificial damage made by electrical discharge machining (EDM), '
  'drilling or engraving, and fourteen with real damage from accelerated lifetime tests, as described in the dataset '
  'publication. KB23, KB24 and KB27, damaged on both rings, were excluded, which leaves 29 bearings and 2319 recordings of '
  '4 s (one file could not be read).', set()),
 ('Every split is defined by bearing identifiers',
  'Every split is defined by bearing identifiers, and all statistics, thresholds and parameters come from the training '
  'bearings. Because the feature design was informed by Paderborn bearings, the Paderborn results are development-stage '
  'evidence. Four protocols are used (Fig. 11). L8, the artificial-to-real protocol of the dataset authors, trains on '
  'K002, KA01, KA05, KA07, KI01, KI05 and KI07 and tests on K001 and ten real-damage bearings, at condition 0 and over all '
  'conditions. L10, their real-damage protocol, splits 15 bearings into three training and two test bearings per class in '
  '10 combinations. A2R trains on K001–K003 and the 12 artificially damaged bearings and tests on the other 14. LOBO '
  'holds out each of the 29 bearings once.', {'14'}),
 ('Fig. 11. Bearing assignment',
  'Fig. 11. Bearing assignment under the four Paderborn protocols; L10 rotates 15 bearings over 10 combinations.', set()),
 ('Table II lists the compared methods',
  'Table II lists the compared methods, whose settings were fixed in advance. Besides the kinematic core, the two branches and a threshold rule, they include WDCNN and a residual network (ResNet-1D) on 4096-sample windows of the envelope or of raw vibration down-sampled to 8 kHz, an impulse-response kernel network whose first layer holds learnable kernels h(t) = exp(−ξωt) sin(ωt), the impulse response of (3), and a fault-order-guided ResNet on the order spectrum (7). Tuned feature baselines and domain-generalization networks (ERM, DANN [41], Deep CORAL [42] and group DRO [43], with training bearings as domains) had their settings selected on the training bearings only.', set()),
 ('Recording accuracy is reported',
  'Recording accuracy is reported, with balanced accuracy, macro F1-score and the macro area under the ROC curve (AUC) '
  'where needed. Because recordings of one bearing are not independent, 95% intervals resample test bearings within each '
  'class (2000 resamples) [44], and methods are compared per bearing with the Wilcoxon signed-rank test (descriptive '
  'p-values). A bearing counts as correct when the class with the highest mean probability is correct. False alarms are '
  'healthy recordings diagnosed as damaged and missed faults damaged recordings diagnosed as healthy, each as a share of '
  'its class; coverage is the share of accepted recordings and selective accuracy the accuracy on them.', set()),
 # ---------------------------------------------------------------- IV
 ('Table III compares all methods',
  'Table III compares all methods. On L8, ReliSense reached 95.9% at condition 0 and 87.0% over all conditions, against '
  '62.3%–65.9% for the published single classifiers, 75.0% for their ensemble and 70.8% for WDCNN; all nine errors at '
  'condition 0 were inner-race recordings diagnosed as outer-race damage (Fig. 12(a)). Its bearing-level 95% interval '
  '(88.6%–100%) lies above the ensemble value, although L8 has only one healthy test bearing. On L10, the published 98.3% '
  'exceeds every result here; Vieira et al. [15] could not reproduce it with bearing-wise splits.', set()),
 ('ReliSense was the most accurate method under all five',
  'ReliSense was the most accurate method under all five protocols, with 78.7% on L10, 77.7% on A2R and 83.7% under LOBO, '
  'but at most 1.5 points above its better branch. All networks were less accurate than ReliSense and its kinematic core '
  'under every protocol. Under LOBO, ReliSense diagnosed 27 of the 29 bearings correctly (95% interval 76.5%–90.6%; macro '
  'AUC 95.7%; balanced accuracy 81.2%; macro F1-score 81.5%) and was more accurate than WDCNN on 16 bearings and less '
  'accurate on 6 (mean difference 15.6 points; Wilcoxon p = 0.009).', set()),
 ('The residual network improved on WDCNN',
  'With raw vibration, the residual and impulse-response kernel networks diagnosed every healthy L8 recording as damaged. The fault-order-guided ResNet was the most accurate of these networks under four protocols (78.8% under LOBO); ReliSense was more accurate than it on 17 bearings and less accurate on 5 (4.9 points; p = 0.038), and than the envelope ResNet on 18 and 3 (17.5 points; p = 0.016). The envelope ResNet raised fewer false alarms than ReliSense with training priors (18.1% against 32.1%) because it missed 28.0% of the faults, against 6.6%; with equal priors, ReliSense had fewer of both (15.6% and 11.2%).', set()),
 ('Baselines whose settings were selected',
  'Baselines tuned on the training bearings did not close the gap: the tuned kinematic core reached 73.2%, 73.5% and '
  '79.7% on L10, A2R and LOBO, and the tuned generic features at most 58.9%. The domain-generalization networks reached at '
  'most 67.7% on L8 at condition 0, 65.8% on L10, 70.4% on A2R and 79.1% under LOBO, with 36.2%–47.7% false alarms under '
  'LOBO.', set()),
 ('L10: mean ± SD over 10 splits',
  'L10: mean ± SD over 10 splits; networks on L8 and A2R: mean ± SD over 3 seeds; WDCNN: 3 seeds, other networks 1 seed. LOBO: recordings of all 29 bearings pooled (WDCNN: mean per-bearing accuracy). Published single classifiers on L8: 62.3–65.9. Tuned baselines not run on L8; Deep CORAL (envelope) on L8 not recorded; group DRO (envelope) not completed.', set()),
 ('The healthy class is the weakest part of the diagnosis outside',
  'The healthy class is the weakest part of the diagnosis outside L8 at condition 0: 65.8% of the healthy test recordings '
  'raised a fault alarm under A2R, 43.8% on L8 over all conditions and 32.1% under LOBO, against 0.6%–6.6% missed faults. '
  'Most LOBO false alarms came from K004 and K005, correct on 25.0% and 31.3% of their recordings and mostly diagnosed as '
  'outer-race damage; their median BPFO feature (0.83 and 0.84) exceeded that of K001–K003 (0.56–0.69).', set()),
 ('Weak excitation lowers both the peaks',
  'Weak excitation lowers both the peaks and the confidence: under LOBO, the median confidence was 0.89 for correct and '
  '0.54 for wrong recordings (Fig. 13(a)). Training-only thresholds for target coverages of 90%, 80% and 70% accepted '
  '88.7%, 78.6% and 68.6% of the recordings, with selective accuracies of 89.9%, 94.1% and 97.5% against 83.7% (Fig. '
  '13(b)), false alarms of 17.5%, 8.1% and 2.7% against 32.1%, and missed faults of 3.4%, 1.5% and 0.5% against 6.6%. At '
  'κ = 0.7, the 95% bearing-level intervals were 57.1%–79.7% for the coverage and 95.1%–99.2% for the selective accuracy, '
  'so the coverage target holds on average, not for every bearing. Acceptance was lowest for K004 and K005 (4% and 5%) and '
  'for the weakest damage (18%–25% for KA30, KA15, KI17 and KA22) (Fig. 13(c)).', set()),
 ('For a comparison at equal coverage',
  'For an equal-coverage comparison, the most confident LOBO recordings of each network were accepted up to the coverage '
  'of ReliSense, a ranking on the test recordings that favours the networks. At 68.6% coverage, ReliSense reached 97.5% '
  'with 2.7% false alarms, or 96.7% with 4.6% under the same ranking, against 63.9%–90.8% with 7.5%–26.9% for the other '
  'networks and 70.2%–89.0% with 13.1%–23.1% for the domain-generalization networks. ReliSense accepted fewer healthy '
  'recordings (45.6% against 61.5%–88.3%).', set()),
 ('Fig. 13. Selective classification',
  'Fig. 13. Selective classification under LOBO. (a) Confidence of correct and incorrect recordings. (b) Accuracy against '
  'coverage; markers: training-only thresholds (κ = 0.9, 0.8, 0.7). (c) Acceptance at κ = 0.7 against bearing accuracy.', set()),
 # ---------------------------------------------------------------- V
 ('To examine whether the kinematic signature',
  'To test whether the kinematic signature transfers beyond one rig and bearing type, the framework was applied without '
  'change to CWRU [45], HUST [46] and the time-varying-speed Ottawa data [47], with segmentation, features, classifiers '
  'and comparators fixed before these data were analyzed.', set()),
 ('On CWRU, one defect diameter was held out',
  'On CWRU, one defect diameter was held out at a time (LOSO), and the single healthy bearing was split by load. On HUST, '
  'one bearing type was held out (LOTO), so every test bearing is new and of an unseen geometry. On Ottawa, one speed '
  'profile was held out (LOPO), which tests only an unseen profile of the same bearings. Each protocol was also run in '
  'reverse, with one diameter, type or profile for training (1SIZE, 1TYPE, 1PROF). The comparators were 32 band energies '
  'with logistic regression, and time and envelope statistics with band energies in a random forest. The main task has '
  'three classes; a four-class task adds ball defects.', set()),
 ('Table V gives the three-class results',
  'Table V gives the three-class results. On CWRU, ReliSense was the most accurate method under both protocols (83.2% and '
  '76.0%). On HUST, where every test bearing is new and of an unseen geometry, it reached 85.0% and, with one training '
  'type, 73.9%, against at most 71.1% and 55.4% for the generic features. The envelope branch raised 60.0% false alarms on '
  'CWRU, whose 12-kHz sampling leaves a band of only 2–5.4 kHz, but averaging kept ReliSense at or above the kinematic '
  'core. On Ottawa, the generic features classified every segment correctly (ReliSense: 98.3% and 95.9%); because each '
  'class is one bearing, features that identify the bearing rather than its fault suffice there.', set()),
 ('The per-bearing results on HUST',
  'Per HUST bearing (Fig. 14), 11 of 15 were diagnosed with at least 90% accuracy under LOTO; the errors of ReliSense '
  'concentrated on the healthy 6205, the outer-race 6208 and the inner-race 6204. With one training type, ReliSense lost '
  '11.1 points and the generic random forest 15.7.', set()),
 ('Fig. 14. Accuracy on each HUST',
  'Fig. 14. Accuracy on each HUST test bearing for ReliSense, its kinematic core and two generic baselines: (a) LOTO; (b) '
  '1TYPE.', set()),
 ('Two weaknesses recurred',
  'ReliSense diagnosed 10.0%–16.7% of the healthy CWRU and HUST segments as faulty, fewer than its kinematic core '
  '(20.0%–36.7%) and, on HUST, the generic methods (up to 45.8%). Ball defects were the weakest class in the four-class '
  'task (recall of the kinematic core 42.2% on CWRU under LOSO and 60.0% on Ottawa), in line with the weak ball signatures '
  'of the CWRU data [45].', set()),
 # ---------------------------------------------------------------- VI
 ('Table VI evaluates the components',
  'Table VI evaluates the components. The 16 modulation features raised the accuracy of the kinematic core by 1.5–4.2 '
  'points under all five protocols. Averaging the two branches gave the highest pooled accuracy under every protocol and, '
  'at κ = 0.7, the highest selective accuracy (97.5% with 2.7% false alarms), equal to the core combined with the envelope '
  'branch.', set()),
 ('Table VII compares the results',
  'Table VII lists published Paderborn results with test bearings absent from training. On L8, ReliSense exceeded the '
  'published classifiers by more than 20 points. Its LOBO macro AUC exceeded those of Vieira et al. [15], whose splits '
  'held only real-damage and healthy bearings, so this comparison is indicative; Perminov and Korzun [48] reported large '
  'differences between folds.', set()),
 ('Evaluation units and class sets differ',
  'Evaluation units and class sets differ; the correspondence of the L10 combinations of this study to the published L10 '
  'evaluation is unknown, so 98.3% is listed for reference only.', set()),
 ('1) What the Two Branches Contribute',
  'The two branches erred on different recordings (Fig. 15(a)): on L8 at condition 0, the kinematic core misclassified twelve inner-race recordings as outer-race damage and the envelope branch 35, the averaged model nine. Averaging mainly made the confidence more reliable: at κ = 0.7, it gave 97.5% selective accuracy with 2.7% false alarms, against 95.7% and 5.0% for the kinematic branch alone. Per bearing, ReliSense was more accurate than the envelope branch on 7 bearings and less accurate on 11 (Wilcoxon p = 0.69), and than the kinematic core on 14 and 5 (p = 0.048; Fig. 15(b)). Over the eleven settings of Table VIII, each branch fell up to 6.3 or 11.4 points below the other, whereas ReliSense was at least as accurate as the better branch in 9 settings and at most 1.6 points below it otherwise: averaging provides robustness across datasets rather than a significant gain on any one. Weighted and stacked combinations and harmonic-agreement features did not improve on equal weights.', {'8', '6.3', '11.4', '9', '1.6'}),  # moved from the Table VIII paragraph of v2.8
 ('Fig. 15. The two branches',
  'Fig. 15. The branches under LOBO. (a) Accuracy of each bearing, kinematic against envelope branch. (b) Accuracy of '
  'ReliSense minus that of the kinematic core per bearing, sorted.', set()),
 ('2) The Healthy Bearings and the Cost',
  'The healthy class was the weakest part of the diagnosis throughout (Fig. 16). The false alarms of K004 and K005 reflect a weak line at the outer-race fault order, which run-in time does not explain (ρ = 0.23, n = 6) and which may come from a surface irregularity, contamination or a rig resonance; these bearings may not be perfectly healthy. At κ = 0.7, ReliSense accepted only 3.8% and 5.0% of their recordings (Fig. 16(b)). Class-wise thresholds shift the balance further, lowering the false alarms of the core with the envelope branch from 2.7% to 0.6% but raising its missed faults from 0.5% to 4.0%.', set()),
 ('The balance can also be shifted before abstention',
  'The balance can also be shifted without retraining. Because the healthy class is the smallest in training (6 of 29 '
  'bearings), dividing each branch posterior by the training class frequencies gives the decision for equal class priors '
  '[49] (Table IX). Under LOBO, this raised the accuracy from 83.7% to 84.7% and lowered the false alarms from 32.1% to '
  '15.6%, while the missed faults rose from 6.6% to 11.2%; K004 and K005 were then correct in 61.3% and 72.5% of their '
  'recordings. Exponents between 0 and 1 on the class frequencies give intermediate operating points. The correction '
  'moves errors between classes: KA15, KA22 and KA30 lost 12.5–16.3 points, and per bearing the two rules did not differ '
  'significantly (p = 0.59). Other results use the training priors unless labelled otherwise.', set()),
 ('LOBO, κ = 0.7: selective accuracy',
  'LOBO, κ = 0.7: selective accuracy/false alarms/missed faults; coverage 68.6% and 69.9%. Balanced accuracy under LOBO: '
  '81.2% and 84.7%. L10, HUST and Ottawa have equal training class frequencies, so both rules agree there.', set()),
 ('Table X places these operating points',
  'Table X compares these operating points with all networks under LOBO. Only the envelope ResNet and the '
  'impulse-response kernel network raised fewer false alarms than ReliSense with training priors (18.1% and 30.4%), and '
  'both were less accurate and missed more faults; with equal priors, ReliSense raised fewer false alarms (15.6%) than '
  'every network.', set()),
 ('ReliSense: training-only thresholds',
  'ReliSense: training-only thresholds (κ = 0.7); with equal priors, 69.9% coverage. Networks: most confident recordings '
  'accepted up to 68.6% coverage, which favours them. WDCNN: mean per-bearing accuracy; missed faults and confidences not '
  'stored.', set()),
 ('Fig. 16. The six healthy',
  'Fig. 16. The six healthy bearings under LOBO. (a) Share of recordings diagnosed as healthy, outer-race or inner-race '
  'damage; FA: false alarms. (b) The same with abstention at κ = 0.7', set()),
 ('3) Relation to Learned and Physics-Informed',
  'None of the compared networks reached ReliSense under any protocol (Table III). The comparisons concern the implemented pipelines, which differed in input information, and do not show that physical priors in learned models fail in general. Only the network that received the fault orders came close, 12.7 points above the same architecture on the envelope, which supports the view that the order domain carries most of what transfers to a new bearing; it remained 4.9 points below ReliSense under LOBO (Fig. 17). The envelope ResNet diagnosed K004, K005 and KI17 without error, so a learned branch restricted to the order domain is a possible extension.', set()),
 ('Fig. 17. Accuracy of each Paderborn',
  'Fig. 17. Accuracy of each bearing under LOBO, ReliSense against the fault-order-guided ResNet; labelled: differences '
  'of 30 points or more.', set()),
 ('4) Deployment',
  'ReliSense needs the bearing geometry, the shaft speed and one recording, but no data of the monitored bearing for '
  'training or calibration, and its computation is small: a few Fourier transforms, a 96-parameter linear model and 500 '
  'trees. The price of reliability is the referral rate: at κ = 0.7, about 31% of the recordings were referred, mostly '
  'from weak damage, low speed and the two healthy bearings that resemble damage.', set()),
 ('On the further datasets (Table XI)',
  'On the further datasets (Table XI), the coverage stayed near its target and the accepted segments were more accurate '
  'where the inner loop could hold out the same kind of unit as the test. With one training fault size or bearing type, '
  'only loads could be held out, whose confidences exceed those of new bearings, so only 18.5% of the CWRU and 27.2% of '
  'the HUST segments were accepted. The calibration therefore needs at least two training bearings per class.', set()),
 ('Training priors. False alarms: healthy segments accepted',
  'Training priors. False alarms: accepted healthy segments diagnosed as faulty, as a share of all healthy segments. With '
  'fewer than two training units of the held-out kind per class, the inner loop holds out loads or recordings.', set()),
 ('5) Decisions From Several Recordings',
  'A monitored bearing is measured repeatedly. Table XII combines consecutive recordings of one bearing and condition by '
  'summing the logarithms of their posteriors, for ReliSense and for the stored posteriors of every network; a sequential '
  'variant adds recordings until the combined confidence reaches 0.95, with at most ten. With training priors, the false '
  'alarms stayed at 30.2%–32.1%, because K004 and K005 were diagnosed as damaged consistently. With equal priors, ten '
  'recordings (40 s) gave 89.6% accuracy with 6.2% false alarms and 8.7% missed faults. The networks reached at most '
  '81.8%, and all except the envelope ResNet (16.7% false alarms, 29.0% missed faults) raised 29.2%–47.9% false alarms. '
  'These rules were evaluated after the main results, and the 6.2% rests on 48 decisions for six healthy bearings, three '
  'of them wrong.', set()),
 ('Posteriors of consecutive recordings',
  'Sum of the log posteriors of consecutive recordings (463 and 231 decisions). Sequential rule: recordings added until '
  'the combined confidence reached 0.95, at most 10; mean 3.2 and 3.3 recordings for ReliSense, 1.5–4.3 for the networks '
  '(1 seed).', set()),
 ('6) Damage Type and Extent',
  'Table XIII groups the LOBO results by the documented damage properties [10]. Artificial damage and real fatigue '
  'pitting, the damage of 9 of the 11 real bearings, were diagnosed with similar accuracy (91.4% and 90.1%). Extent levels '
  '2 and 3 reached 93.6%, against 84.1% for level 1. The weakest damaged bearings, KA15 (68.8%) and KA30 (42.5%), carry '
  'plastic deformation by indentations; KA30, the only single-damage bearing documented as distributed, does not produce '
  'the single impact per passage assumed by (3). Without KA30, LOBO accuracy was 85.2%, and 98.1% at 70.4% coverage; all '
  'other results include it, because the damage type of a monitored bearing is unknown.', set()),
 ('Damage properties from the dataset documentation',
  'Damage properties from [10]; extent levels as in Table I. Equal priors: Section VI-C, part 2. κ = 0.7: training '
  'priors.', set()),
 ('7) Changes Tested Under an Adoption Rule',
  'After the main results, four feature changes were tested under a rule fixed before the runs: the change with the highest mean balanced accuracy over the eleven settings would replace ReliSense only if it was not less accurate under LOBO and lost at most 0.5 points of selective accuracy at κ = 0.7. The changes were condition-wise standardization, a multiband kinematic branch, and comb-separated and slip-consistent features that separate the fault lines from the integer shaft orders in the fault windows of the 6203 bearing [2]. No change met the rule (Table XIV). The equal-prior decision had the highest mean (85.5% against 83.9%) but lowered the selective accuracy by 1.1 points; the slip-consistent features helped where test bearings differed in size or type from training but lost 5.0 points on L8 at condition 0.', set()),
 ('Mean over the eleven settings of Table VIII',
  'Mean over the eleven settings of Table VIII. Changes replace the kinematic branch; condition-wise standardization was '
  'applied to both branches. Last column: coverage 68.3%–69.9%.', set()),
 # ---------------------------------------------------------------- VII, VIII, data
 ('The study has limitations.',
  'The study has limitations. The main analysis uses one rig, one bearing type and four operating conditions, and the real damage grew in accelerated lifetime tests rather than in service. The further datasets have seeded damage and few bearings, and no model was tested across machines. Fifteen of the 29 bearings informed the design of the kinematic core, and the modulation features and envelope branch were added after its Paderborn results were known, so these results are development-stage evidence. The networks did not receive the bandwidth and harmonic range of ReliSense, the physics-informed networks followed published principles rather than original code, and several were trained with one seed. The equal-prior decision, the decisions from several recordings and the changes in Section VI-C, part 7, were evaluated after the main results. Tests across machines, for example on the 20 bearings of the Ottawa constant-speed collection [50], calibration under distribution shift [51], [52] and conformal prediction [53] are further directions.', set()),
 ('This article addressed the diagnosis',
  'ReliSense diagnoses unseen bearings by reading the squared envelope spectrum at the fault orders fixed by geometry and measured speed, classifying these features in two branches and withholding low-confidence diagnoses with thresholds set on training bearings only. On 29 Paderborn bearings, which also served for method development, it reached 95.9% when trained on artificial and tested on real damage, against 75.0% for the published ensemble and at most 70.8% for the compared networks. Under leave-one-bearing-out evaluation, it reached 83.7%, and 97.5% at 69% coverage with 2.7% false alarms. Equal class priors lowered the false alarms before abstention to 15.6%, and ten consecutive recordings to 6.2%, against at least 16.7% for every compared network. It was also the most accurate method on unseen bearing types of HUST and fault sizes of CWRU. Healthy bearings that resemble damage, weak or distributed damage and low speed remain its main limits.', set()),
 ('The Paderborn data are available',
  'The Paderborn data are available from the KAt-DataCenter of Paderborn University (CC BY-NC 4.0), and the CWRU, HUST '
  'and Ottawa data from their providers. The code, label sheet, split manifests, the identifier of the unreadable '
  'recording and the per-fold predictions will be released at [repository link to be added].', set()),
 ('The results support one central finding',
  'The central finding is that a diagnosis in the coordinates fixed by the bearing kinematics carries over to unseen bearings.', set()),
 ('Segments pooled over folds. False alarms',
  'Segments pooled over folds. On CWRU the healthy bearing also supplies training data, on Ottawa every test bearing does. Kinematic core false alarms: 36.7, 20.0, 21.7, 20.8, 2.3 and 5.4.', set()),
 ('Paderborn: recordings; further datasets',
  'Paderborn: recordings; further datasets: segments; L10: mean over 10 splits. Equal priors: Section VI-C, part 2.', set()),
]
SENTENCE_EDITS = [   # paragraphs with formatted symbols: only plain parts are shortened
 ('Dividing by fr gives the fault orders', 'Fig. 2(a) shows the rolling kinematics. Fig. 2(c) shows the 15 search windows used in Module I.',
  'Fig. 2 shows the kinematics and the 15 search windows of Module I.'),
 ('Dividing by fr gives the fault orders',
  'A local defect can lie on the outer raceway, the inner raceway or a rolling element (Fig. 3). Each passage of a rolling '
  'element over the defect produces a short impact that excites a structural resonance, so the vibration measured on the '
  'housing can be modelled as',
  'For a local defect on a raceway or rolling element (Fig. 3), each passage produces an impact that excites a structural '
  'resonance, so the housing vibration can be modelled as'),
 ('where θi = 2π·fr·i·Td', ', independently of how the defect was produced and of the mounting or sensor', ''),
 ('1) Band Selection and Demodulation', 'The impacts of (3) are visible as a repetition rate only after demodulation. ', ''),
 ('5) Modulation Features', 'An inner-race defect is modulated by the load zone at the shaft rate and a ball defect by the '
  'cage, which produces sidebands around the fault lines. ', ''),
 ('6) Envelope Features', ' No band energy or other generic spectral feature is used.', ''),
 ('where τb is the threshold for test bearing b', ', and it should reflect the lower confidence that a model produces on a '
  'bearing it has never seen. It is therefore calibrated', '; it is calibrated'),
 ('Table IV summarizes the datasets', 'All features were read at the fault orders of each bearing, with the HUST geometry '
  'taken from its data description.', 'The HUST geometry was taken from its data description.'),
 ('The search band absorbs speed errors', ' A healthy bearing gives φkh close to zero, and a defect raises the features of its own fault order.', ''),
 ('where Q1−κ is the empirical', ' With |Tb| = 28 under LOBO, each threshold requires 28 inner fits of both branches.', ''),
 ('For a 4-s recording, the frequency resolution', ' For an inner-race defect, it also contains lines at h·fd ± m·fr.', ''),
 ('Table IV summarizes the datasets', 'Every recording was cut into segments of 50 revolutions, and the squared envelope was resampled to the shaft angle.', 'The squared envelope was resampled to the shaft angle.'),
]
TABLE2 = {   # Table II (index 1): (row, column) -> shorter text
 (1, 1): '31 kinematic and 62 envelope features',
 (1, 2): 'Kinematic branch: LR, L2, C = 0.3 (96 parameters); envelope branch: RF, 500 trees; posteriors averaged; thresholds by inner LOBO',
 (2, 2): 'LR, C = 0.3; 48 parameters (first design stage)',
 (3, 2): 'RF, 500 trees, seed 0',
 (4, 2): '95th percentile of healthy training recordings',
 (6, 2): 'Four stages of two residual blocks (32–128 channels); 1 020 195 parameters',
 (7, 2): '32 damped-resonance kernels (8 ms) and a WDCNN-type stack; 86 195 parameters',
 (9, 2): 'LR, RBF-SVM, RF or gradient boosting; classifier and hyperparameters by stratified group K-fold',
 (10, 1): 'Envelope windows (WDCNN backbone) or order spectrum (1-D CNN)',
 (10, 2): 'ERM, DANN [41], Deep CORAL [42], group DRO [43]; domains = training bearings; 8 bearings × 16 per batch; strength chosen on a bearing-wise validation split',
}
DELETE = ['Table VIII compares the two branches']   # its content is merged into part 1
DROP_COLUMN = (11, 1)   # Table XII: the single-recording column repeats Table X
COMPACT = re.compile(r'(?<=[\d–])\s/\s(?=[\d–])')


# ------------------------------------------------------------------ helpers
def all_pars(d):
    ps = list(d.paragraphs)
    for t in d.tables:
        for r in t.rows:
            for c in r.cells: ps += c.paragraphs
    return ps


def find(d, start):
    hits = [p for p in d.paragraphs if p.text.startswith(start)]
    assert len(hits) == 1, (start, len(hits)); return hits[0]


def nums(t): return set(re.findall(r'\d+(?:\.\d+)?', t))


def has_head(p):
    rs = p.runs
    return len(rs) > 1 and rs[0].italic and re.match(r'\d\) ', rs[0].text) is not None


def rewrite(p, new):
    rs = list(p.runs); head = rs[0] if has_head(p) else None
    body = [r for r in rs if r is not head]; keep = max(body, key=lambda r: len(r.text))
    for r in body:
        if r is not keep: r._r.getparent().remove(r._r)
    for tag in ('hyperlink', 'proofErr'):
        for x in p._p.findall(W + tag): p._p.remove(x)
    keep.text = new


def replace_across_runs(p, old, new):
    runs = p.runs; full = ''.join(r.text for r in runs); k = full.find(old)
    assert k >= 0 and full.count(old) == 1, (old[:70], full[:90])
    pos = 0; first = True; end = k + len(old)
    for r in runs:
        a, b = pos, pos + len(r.text); pos = b
        if b <= k or a >= end: continue
        s, e = max(k, a) - a, min(end, b) - a
        if first: r.text = r.text[:s] + new + r.text[e:]; first = False
        else: r.text = r.text[:s] + r.text[e:]


def set_cell(cell, text):
    ps = cell.paragraphs
    for extra in ps[1:]: extra._p.getparent().remove(extra._p)
    rs = ps[0].runs; rs[0].text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)


def drop_column(tbl, j):
    t = tbl._tbl; grid = t.find(W + 'tblGrid'); cols = grid.findall(W + 'gridCol')
    total = sum(int(c.get(W + 'w')) for c in cols); grid.remove(cols[j]); rest = grid.findall(W + 'gridCol')
    scale = total / sum(int(c.get(W + 'w')) for c in rest)
    ws = [int(int(c.get(W + 'w')) * scale) for c in rest]
    for c, w in zip(rest, ws): c.set(W + 'w', str(w))
    for tr in t.findall(W + 'tr'):
        tcs = tr.findall(W + 'tc'); tr.remove(tcs[j])
        for tc, w in zip(tr.findall(W + 'tc'), ws):
            tcw = tc.find(W + 'tcPr/' + W + 'tcW')
            if tcw is not None: tcw.set(W + 'w', str(w))


def mark(p, oldtext):
    new = p.text; runs = list(p.runs); head = runs[0] if has_head(p) and oldtext.startswith(runs[0].text) else None
    off = len(head.text) if head is not None else 0
    A = re.findall(r'\S+\s*', oldtext[off:]); B = re.findall(r'\S+\s*', new[off:])
    tmpl = copy.deepcopy(max([r for r in runs if r is not head], key=lambda r: len(r.text))._r)
    for r in runs:
        if r is not head: r._r.getparent().remove(r._r)
    def add(t, col=None, strike=False):
        r = p.add_run(t); rpr = tmpl.find(W + 'rPr')
        if rpr is not None: r._r.insert(0, copy.deepcopy(rpr))
        if col: r.font.color.rgb = col
        if strike: r.font.strike = True
    for op, a1, a2, b1, b2 in difflib.SequenceMatcher(None, A, B, autojunk=False).get_opcodes():
        if op == 'equal': add(''.join(A[a1:a2]))
        if op in ('delete', 'replace'): add(''.join(A[a1:a2]), RED, True)
        if op in ('insert', 'replace'): add(''.join(B[b1:b2]), BLUE)


def word_count(path):
    d = Document(path); n = sum(len(p.text.split()) for p in d.paragraphs)
    return n, n + sum(len(''.join(x.text or '' for x in tc.iter(W + 't')).split()) for t in d.tables for tc in t._tbl.iter(W + 'tc'))


def build(marked):
    d = Document(src); body_pars = {p._p: p.text for p in d.paragraphs}
    for start, new, allowed in REWRITE:
        p = find(d, start); o = p.text
        body_old = o[len(p.runs[0].text):] if has_head(p) else o
        extra = nums(new) - nums(body_old) - allowed
        assert not extra, (start, sorted(extra))
        rewrite(p, new)
    for start, o, n in SENTENCE_EDITS: replace_across_runs(find(d, start), o, n)
    for start in DELETE:
        p = find(d, start)
        if marked:
            for r in p.runs: r.font.color.rgb = RED; r.font.strike = True
        else: p._p.getparent().remove(p._p)
    t2 = d.tables[1]
    for (i, j), txt in TABLE2.items(): set_cell(t2.rows[i].cells[j], txt)
    assert d.tables[DROP_COLUMN[0]].rows[0].cells[DROP_COLUMN[1]].text == '1 recording'
    drop_column(d.tables[DROP_COLUMN[0]], DROP_COLUMN[1])
    for t in d.tables:
        for tc in t._tbl.iter(W + 'tc'):
            for r in tc.iter(W + 'r'):
                for x in r.iter(W + 't'):
                    if x.text and COMPACT.search(x.text): x.text = COMPACT.sub('/', x.text)
    if marked:
        for p in d.paragraphs:
            if p._p in body_pars and p.text != body_pars[p._p] and not any(p.text.startswith(x) for x in DELETE): mark(p, body_pars[p._p])
        P = d.paragraphs; k = [i for i, p in enumerate(P) if p.text.startswith('Revision key')]
        if k: P[k[0]]._p.getparent().remove(P[k[0]]._p)
        key = d.paragraphs[0].insert_paragraph_before('')
        key.add_run('Revision key (v2.8 → v2.9, text condensed; tables: same values in compact notation, Table II shortened, '
                    'Table XII without its single-recording column):  ')
        x = key.add_run('red strikethrough = deletion'); x.font.color.rgb = RED; x.font.strike = True
        key.add_run('   |   '); x = key.add_run('blue = addition'); x.font.color.rgb = BLUE
    return d


build(False).save(out_clean); build(True).save(out_marked)
n, m = word_count(out_clean); print(f'words: paragraphs {n}; paragraphs + tables {m}')
