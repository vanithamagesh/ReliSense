"""Paper A v2.6 from v2.5: tuned feature baselines and domain-generalization networks (phase 24), false alarms read
together with missed faults for all methods (phases 23-26, new Table X), references [41]-[43] (DANN, Deep CORAL,
group DRO; old [41]-[50] -> [44]-[53]).  python3 make_A26.py <v2_5.docx> <out_clean.docx> <out_marked.docx>
Real numbers only: phase24_tables_pasted.txt and the pasted part 2 tables, phase25_matched_coverage_pasted.txt,
phase26_output_pasted.txt, phase26b_output_pasted.txt; WDCNN false alarms from wdcnn_env_lobo.json (1 - mean healthy
recall = 42.9%). Marked copy: red strikethrough = deletion, blue = addition."""
import sys, copy, re, difflib
from docx import Document
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph
from docx.table import Table, _Row

src, out_clean, out_marked = sys.argv[1:4]
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
RED, BLUE = RGBColor(255, 0, 0), RGBColor(0, 0, 255)

EDITS = [   # (anchor contained in the paragraph, old, new) - applied after the citation renumbering
 ('Paper A, version', 'Paper A, version 2.5, 8 October 2026', 'Paper A, version 2.6, 8 October 2026'),
 ('The published vibration ensemble reached', 'and convolutional, residual and physics-informed networks at most 70.8%',
  'and convolutional, residual, physics-informed and domain-generalization networks at most 70.8%'),
 ('The published vibration ensemble reached', 'halved the false alarms before abstention to 15.6%, at the cost of more missed faults (11.2% instead of 6.6%).',
  'lowered the false alarms before abstention to 15.6% with 11.2% missed faults, against 18.1% and 28.0% for the envelope residual network.'),
 ('as in fault-frequency-guided networks', 'as in fault-frequency-guided networks.',
  'as in fault-frequency-guided networks. Two further groups had their settings selected on the training bearings only. '
  'The tuned feature baselines chose a logistic regression, an RBF support vector machine, a random forest or gradient '
  'boosting, together with its hyperparameters, by stratified group K-fold over the training bearings (groups = bearings), '
  'for the band energies, the time and band statistics and the kinematic core; L8 was not evaluated, because its single '
  'healthy training bearing cannot be held out. The domain-generalization networks treat every training bearing as a domain '
  'and were trained with class-weighted cross-entropy (ERM), domain-adversarial training (DANN) [41], Deep CORAL [42] or '
  'group DRO [43], on envelope windows with the WDCNN backbone or on the order spectrum with a small one-dimensional '
  'network. Each batch held 16 windows or recordings from each of eight bearings, and the strength of the domain term was '
  'chosen on a validation split that held out one training bearing per class.'),
 ('under LOBO than ReliSense (18.1%',
  'The envelope ResNet raised fewer false alarms under LOBO than ReliSense (18.1% against 32.1%), but it missed 28.0% of the faults, against 6.6%, and on L8 it diagnosed',
  'Under LOBO, the envelope ResNet raised fewer false alarms than ReliSense with training priors (18.1% against 32.1%) '
  'because it diagnosed many damaged recordings as healthy: it missed 28.0% of the faults, against 6.6%. With equal class '
  'priors (Section VI-C), ReliSense raised fewer false alarms (15.6%) and missed fewer faults (11.2%) than this network. '
  'On L8, the envelope ResNet diagnosed'),
 ('L10: mean ± SD over 10 splits', 'Published single classifiers on L8: 62.3–65.9.',
  'Published single classifiers on L8: 62.3–65.9. Tuned baselines (classifier and hyperparameters chosen on the training '
  'bearings) and domain-generalization networks: L10 pooled over the 10 splits, 1 seed; tuned baselines not evaluated on L8; '
  'L8 results of Deep CORAL on envelope input not recorded; group DRO on envelope input not completed. Equal priors: Section VI-C.'),
 ('To compare abstention at equal coverage', 'against 61.5%–88.3% for these networks.',
  'against 61.5%–88.3% for these networks. Under the same ranking, the domain-generalization networks reached 70.2%–89.0% '
  'selective accuracy with 13.1%–23.1% false alarms (Table X).'),
 ('These comparisons concern the implemented pipelines',
  'These comparisons concern the implemented pipelines, which differed in input information and were not tuned, and do not show',
  'These comparisons concern the implemented pipelines, which differed in input information. The networks of Section IV-A '
  'were not tuned, but the tuned feature baselines and the domain-generalization networks, whose settings were chosen on the '
  'training bearings, did not reach ReliSense either. The comparisons do not show'),
 ('Abstention was also applied to the further', '(Table X)', '(Table XI)'),
 ('were not tuned. The equal-coverage', 'code and were not tuned. ',
  'code and were not tuned. The domain-generalization networks were trained with one seed, and two of their configurations are incomplete. '),
 ('at most 70.8% for the compared learned', 'at most 70.8% for the compared learned and physics-informed networks',
  'at most 70.8% for the compared learned, physics-informed and domain-generalization networks'),
 ('at most 70.8% for the compared learned', 'halved the false alarms before abstention, at the cost of more missed faults.',
  'lowered the false alarms before abstention to 15.6% with 11.2% missed faults, fewer of both than the envelope residual network (18.1% and 28.0%).'),
]
CAPTION_EDITS = [('TABLE X', 'TABLE XI')]   # exact paragraph text replacement (the external-abstention table)

NEW_PAR_AFTER_ANCHOR = {
 'under LOBO than ReliSense (18.1%':
  "Baselines whose settings were selected on the training bearings did not close the gap (Table III). With the classifier "
  "and its hyperparameters chosen by grouped cross-validation, the band energies reached 50.5%–58.9%, the time and band "
  "statistics 45.7%–55.5% and the kinematic core 73.2%, 73.5% and 79.7% on L10, A2R and LOBO, against 78.7%, 77.7% and "
  "83.7% for ReliSense with its fixed settings. The domain-generalization networks reached at most 67.7% on L8 at "
  "condition 0, 65.8% on L10, 70.4% on A2R and 79.1% under LOBO. Deep CORAL improved the order-spectrum network under all "
  "five protocols, by 1.0–5.9 points, whereas DANN gave the highest LOBO accuracy of these networks (79.1%) but fell to "
  "24.1% on L8 at condition 0. Under LOBO, all domain-generalization networks raised more false alarms than ReliSense "
  "(36.2%–47.7%).",
 'The balance can also be shifted':
  "Table X places these operating points among all methods evaluated under LOBO. Before abstention, the domain-"
  "generalization networks raised 36.2%–47.7% false alarms and the other networks 18.1%–45.4%. Only the envelope ResNet "
  "(18.1%) and the impulse-response kernel network (30.4%) raised fewer false alarms than ReliSense with training priors, "
  "and both were less accurate (66.1% and 62.1%) and missed more faults (28.0% and 8.7%). ReliSense with equal priors "
  "raised fewer false alarms than every network in Table X (15.6%) and missed fewer faults than the envelope ResNet. With "
  "abstention, the difference grew: at the coverage of ReliSense, 68.6%, the networks still raised 7.5%–26.9% false alarms, "
  "against 2.7% for ReliSense with its training-only thresholds and 4.6% under the same confidence ranking as the networks.",
}
NEW_TABLE_AFTER_ANCHOR = {
 'The balance can also be shifted': ('TABLE X', 'Accuracy, False Alarms and Missed Faults of All Methods Under LOBO (%)',
  ['Method', 'Without abstention: accuracy / false alarms / missed faults', 'At 68.6% coverage: selective accuracy / false alarms / missed faults'],
  [['ReliSense, training priors', '83.7 / 32.1 / 6.6', '97.5 / 2.7 / 0.5'],
   ['ReliSense, prior exponent 0.5', '84.3 / 22.7 / 9.1', '–'],
   ['ReliSense, equal priors', '84.7 / 15.6 / 11.2', '96.4 / 1.7 / 1.9'],
   ['Fault-order-guided ResNet', '78.8 / 37.5 / 8.3', '90.8 / 12.1 / 2.5'],
   ['ResNet-1D, envelope', '66.1 / 18.1 / 28.0', '75.0 / 7.5 / 17.9'],
   ['ResNet-1D, raw vibration', '59.5 / 45.4 / 10.6', '63.9 / 26.9 / 6.8'],
   ['Impulse-response kernel network', '62.1 / 30.4 / 8.7', '65.4 / 21.5 / 2.9'],
   ['WDCNN, envelope', '68.1 / 42.9 / –', '–'],
   ['ERM, envelope', '75.9 / 36.9 / 5.5', '86.3 / 19.0 / 1.9'],
   ['DANN, envelope', '59.1 / 41.7 / 9.2', '70.2 / 21.0 / 4.3'],
   ['Deep CORAL, envelope', '72.7 / 47.7 / 2.7', '79.3 / 23.1 / 1.0'],
   ['ERM, order spectrum', '75.8 / 39.0 / 8.5', '89.0 / 15.4 / 2.4'],
   ['DANN, order spectrum', '79.1 / 36.2 / 4.8', '87.0 / 17.3 / 2.0'],
   ['Deep CORAL, order spectrum', '76.8 / 37.3 / 9.6', '88.6 / 13.1 / 3.7'],
   ['Group DRO, order spectrum', '75.9 / 38.8 / 8.6', '88.4 / 15.4 / 2.6']],
  'ReliSense at 68.6% coverage: training-only thresholds (κ = 0.7); with equal priors the same thresholds accepted 69.9%. '
  'Networks: the most confident recordings accepted up to 68.6% coverage, a ranking on the test recordings that favours them. '
  'Missed faults: damaged recordings diagnosed as healthy, as a share of all damaged recordings. WDCNN: mean per-bearing '
  'accuracy; false alarms from its six healthy bearings; its missed faults and per-recording confidences were not stored.'),
}
TABLE_ROWS = {   # table index in v2.5 -> (prototype row index, insert after row index, rows)
 2: [(2, 1, [['ReliSense, equal priors', '95.5', '88.6', '78.7 ± 11.8', '77.9', '84.7']]),
     (11, 11, [['Tuned band energies', '–', '–', '50.5', '58.9', '52.1'],
               ['Tuned time and band statistics', '–', '–', '45.7', '50.7', '55.5'],
               ['Tuned kinematic core', '–', '–', '73.2', '73.5', '79.7'],
               ['ERM, envelope', '67.7', '61.1', '58.9', '62.1', '75.9'],
               ['DANN, envelope', '56.8', '44.3', '52.2', '65.5', '59.1'],
               ['Deep CORAL, envelope', '–', '–', '60.8', '69.9', '72.7'],
               ['ERM, order spectrum', '63.2', '59.0', '63.1', '67.1', '75.8'],
               ['DANN, order spectrum', '24.1', '60.2', '50.3', '66.2', '79.1'],
               ['Deep CORAL, order spectrum', '66.4', '64.9', '65.8', '70.4', '76.8'],
               ['Group DRO, order spectrum', '63.2', '58.5', '63.1', '63.4', '75.9']])],
 1: [(8, 8, [['Tuned feature baselines', 'Band energies; time and band statistics; kinematic core (15)',
              'LR, RBF-SVM, RF (300 trees) or gradient boosting (200 iterations); classifier and hyperparameters by stratified group K-fold over the training bearings'],
             ['Domain-generalization networks', 'Envelope windows (WDCNN backbone) or order spectrum up to 11.8 orders (1-D CNN)',
              'ERM, DANN [41], Deep CORAL [42], group DRO [43]; domains = training bearings; 8 bearings × 16 per batch; AdamW, 20 / 60 epochs; strength chosen on a bearing-wise validation split']])],
}
NEW_REFS_AFTER = '[40]\t'
NEW_REFS = ['[41]\tY. Ganin, E. Ustinova, H. Ajakan, P. Germain, H. Larochelle, F. Laviolette, M. Marchand, and V. Lempitsky, '
            '“Domain-adversarial training of neural networks,” J. Mach. Learn. Res., vol. 17, no. 59, pp. 1–35, 2016.',
            '[42]\tB. Sun and K. Saenko, “Deep CORAL: Correlation alignment for deep domain adaptation,” in Proc. Eur. Conf. '
            'Comput. Vis. Workshops, 2016, pp. 443–450.',
            '[43]\tS. Sagawa, P. W. Koh, T. B. Hashimoto, and P. Liang, “Distributionally robust neural networks for group shifts: '
            'On the importance of regularization for worst-case generalization,” in Proc. Int. Conf. Learn. Represent. (ICLR), 2020.']
SHIFT_FROM, SHIFT = 41, 3


def all_pars(d):
    ps = list(d.paragraphs)
    for t in d.tables:
        for r in t.rows:
            for c in r.cells: ps += c.paragraphs
    return ps


def renumber(d):
    def f(m):
        n = int(m.group(1)); return f'[{n + SHIFT}]' if n >= SHIFT_FROM else m.group(0)
    for p in all_pars(d):
        full = p.text
        for r in p.runs:
            if '[' in r.text: r.text = re.sub(r'\[(\d+)\]', f, r.text)
        assert len(re.findall(r'\[\d+\]', full)) == len(re.findall(r'\[\d+\]', p.text)), full[:80]


def replace_across_runs(p, old, new):
    runs = p.runs; full = ''.join(r.text for r in runs); k = full.find(old)
    assert k >= 0 and full.count(old) == 1, (old[:70], full[:90])
    pos = 0; first = True; end = k + len(old)
    for r in runs:
        a, b = pos, pos + len(r.text); pos = b
        if b <= k or a >= end: continue
        s, e = max(k, a) - a, min(end, b) - a
        if first: r.text = r.text[:s] + new + r.text[e:]; first = False
        else: r.text = r.text[:s] + r.text[e:]


def find(d, anchor, exact=False):
    hits = [p for p in d.paragraphs if (p.text == anchor if exact else anchor in p.text)]
    assert len(hits) == 1, (anchor, len(hits)); return hits[0]


def new_par_after(anchor_el, template_p, text):
    el = copy.deepcopy(template_p._p)
    rs = el.findall(W + 'r'); keep = max(rs, key=lambda r: len(''.join(t.text or '' for t in r.iter(W + 't'))))
    for r in rs:
        if r is not keep: el.remove(r)
    for tag in ('hyperlink', 'bookmarkStart', 'bookmarkEnd', 'proofErr'):
        for x in el.findall(W + tag): el.remove(x)
    anchor_el.addnext(el); p = Paragraph(el, template_p._parent); p.runs[0].text = text
    return p


def setc(cell, text):
    ps = cell.paragraphs
    for extra in ps[1:]: extra._p.getparent().remove(extra._p)
    p = ps[0]; rs = p.runs
    if not rs: p.add_run(text); return
    rs[0].text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)


def build_table(tmpl_tbl, header, rows):
    t = copy.deepcopy(tmpl_tbl._tbl); n = len(header)
    grid = t.find(W + 'tblGrid'); cols = grid.findall(W + 'gridCol')
    total = sum(int(c.get(W + 'w')) for c in cols); w = total // n
    for c in cols[n:]: grid.remove(c)
    while len(grid.findall(W + 'gridCol')) < n: grid.append(copy.deepcopy(cols[0]))
    for c in grid.findall(W + 'gridCol'): c.set(W + 'w', str(w))
    trs = t.findall(W + 'tr'); hdr, body = trs[0], trs[1]
    for tr in trs[2:]: t.remove(tr)
    for tr in (hdr, body):
        for tc in tr.findall(W + 'tc')[n:]: tr.remove(tc)
        while len(tr.findall(W + 'tc')) < n: tr.append(copy.deepcopy(tr.findall(W + 'tc')[-1]))
        for tc in tr.findall(W + 'tc'):
            tcw = tc.find(W + 'tcPr/' + W + 'tcW')
            if tcw is not None: tcw.set(W + 'w', str(w)); tcw.set(W + 'type', 'dxa')
    t.remove(body)
    for _ in rows: t.append(copy.deepcopy(body))
    table = Table(t, tmpl_tbl._parent)
    for j, h in enumerate(header): setc(table.rows[0].cells[j], h)
    for i, row in enumerate(rows):
        for j, v in enumerate(row): setc(table.rows[i + 1].cells[j], v)
    return table


def mark(p, oldtext):
    new = p.text; A = re.findall(r'\S+\s*', oldtext); B = re.findall(r'\S+\s*', new)
    tmpl = copy.deepcopy(max(p.runs, key=lambda r: len(r.text))._r)
    for r in list(p.runs): r._r.getparent().remove(r._r)
    def add(t, col=None, strike=False):
        r = p.add_run(t); rpr = tmpl.find(W + 'rPr')
        if rpr is not None: r._r.insert(0, copy.deepcopy(rpr))
        if col: r.font.color.rgb = col
        if strike: r.font.strike = True
    for op, a1, a2, b1, b2 in difflib.SequenceMatcher(None, A, B, autojunk=False).get_opcodes():
        if op == 'equal': add(''.join(A[a1:a2]))
        if op in ('delete', 'replace'): add(''.join(A[a1:a2]), RED, True)
        if op in ('insert', 'replace'): add(''.join(B[b1:b2]), BLUE)


def blue(p):
    for r in p.runs: r.font.color.rgb = BLUE


def build(marked):
    d = Document(src)
    old = {p._p: p.text for p in all_pars(d)}
    renumber(d)
    anchors = {a: find(d, a) for a in set(list(NEW_PAR_AFTER_ANCHOR) + list(NEW_TABLE_AFTER_ANCHOR))}
    for a, o, n in EDITS: replace_across_runs(find(d, a), o, n)
    for o, n in CAPTION_EDITS: replace_across_runs(find(d, o, exact=True), o, n)
    cap_no = find(d, 'TABLE VI', exact=True)
    P = d.paragraphs; i6 = [k for k, p in enumerate(P) if p.text == 'TABLE VI'][0]; cap_title = P[i6 + 1]
    foot = find(d, 'Last column: selective accuracy / false alarms with training-only thresholds')
    tmpl_tbl = d.tables[5]
    for t_idx, specs in TABLE_ROWS.items():
        tb = d.tables[t_idx]
        for proto, after, rows in sorted(specs, key=lambda s_: -s_[1]):          # bottom first: indices stay valid
            proto_tr = tb.rows[proto]._tr; anchor = tb.rows[after]._tr
            for row in rows:
                tr = copy.deepcopy(proto_tr); anchor.addnext(tr); anchor = tr
                cells = _Row(tr, tb).cells
                for j, v in enumerate(row): setc(cells[j], v)
    POS = {'The balance can also be shifted': find(d, 'LOBO, κ = 0.7: selective accuracy / false alarms / missed faults with training-only')}
    for a, txt in NEW_PAR_AFTER_ANCHOR.items():
        p0 = anchors[a]; pos = POS.get(a, p0); q = new_par_after(pos._p, p0, txt)
        if a in NEW_TABLE_AFTER_ANCHOR:
            no, title, hdr, rows, fn = NEW_TABLE_AFTER_ANCHOR[a]
            q1 = new_par_after(q._p, cap_no, no); q2 = new_par_after(q1._p, cap_title, title)
            tb = build_table(tmpl_tbl, hdr, rows); q2._p.addnext(tb._tbl); new_par_after(tb._tbl, foot, fn)
    r40 = [p for p in d.paragraphs if p.text.startswith(NEW_REFS_AFTER)][0]; anchor = r40
    for ref in NEW_REFS: anchor = new_par_after(anchor._p, r40, ref)
    if marked:
        for p in all_pars(d):
            if p._p not in old: blue(p)
            elif p.text != old[p._p]: mark(p, old[p._p])
        P = d.paragraphs; k = [i for i, p in enumerate(P) if p.text.startswith('Revision key')]
        if k: P[k[0]]._p.getparent().remove(P[k[0]]._p)
        key = d.paragraphs[0].insert_paragraph_before('')
        key.add_run('Revision key (v2.5 → v2.6):  ')
        x = key.add_run('red strikethrough = deletion'); x.font.color.rgb = RED; x.font.strike = True
        key.add_run('   |   '); x = key.add_run('blue = addition'); x.font.color.rgb = BLUE
        key.add_run('   |   Reference numbers from [41] onward increased by 3.')
    return d


build(False).save(out_clean); build(True).save(out_marked); print('ok')
