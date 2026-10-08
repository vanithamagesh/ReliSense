"""Paper A v2.7 from v2.6: decisions from several recordings for ReliSense and all networks (phase 28, new Table XII),
LOBO results by documented damage properties (phase 22 posteriors, label sheet, phase 28 per-bearing results, new
Table XIII), changes tested under an adoption rule fixed in advance (phases 27 and 28, new Table XIV), and the
corresponding edits to the abstract, Section VI-C, the limitations and the conclusion; Wilcoxon p against the envelope branch
from phase 28 (0.69; v2.6: 0.73, same 7 / 11 bearings); Table X no longer cited first in Section IV-C. No new references.
  python3 make_A27.py <v2_6.docx> <out_clean.docx> <out_marked.docx>
Real numbers only. Every table value is read from phase28_results.json, phase27_results.json (condition-wise
standardization of both branches; the phase 28 run may mix two settings for these fusions) and, inside
ReliSense_v3_26_figures_sources.zip, figure_data/phase22_figdata.npz and figure_data/labels_32.csv.
Marked copy: red strikethrough = deletion, blue = addition."""
import sys, copy, re, csv, io, json, zipfile, difflib
import numpy as np
from docx import Document
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph
from docx.table import Table

src, out_clean, out_marked = sys.argv[1:4]
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
RED, BLUE = RGBColor(255, 0, 0), RGBColor(0, 0, 255)
def f1(x): return f'{100 * x:.1f}'

# ------------------------------------------------------------------ numbers
R28 = json.load(open('phase28_results.json')); R27 = json.load(open('phase27_results.json'))
L28 = R28['paderborn']['lobo']; NET = R28['paderborn']['networks_seq']
def sets(r): s = dict(r['paderborn']['settings']); s.update(r['external']); return s
S28, S27 = sets(R28), sets(R27); NAMES = list(S28); assert len(NAMES) == 11 and list(S27) == NAMES

def seqcell(q): return f'{f1(q["acc"])} / {f1(q["fa"])} / {f1(q["missed"])}'
NETS = [('p23_fo_resnet', 'Fault-order-guided ResNet'), ('p23_resnet_env', 'ResNet-1D, envelope'),
        ('p23_resnet_raw', 'ResNet-1D, raw vibration'), ('p23_irk_raw', 'Impulse-response kernel network'),
        ('p24_env_erm', 'ERM, envelope'), ('p24_env_dann', 'DANN, envelope'), ('p24_env_coral', 'Deep CORAL, envelope'),
        ('p24_ord_erm', 'ERM, order spectrum'), ('p24_ord_dann', 'DANN, order spectrum'),
        ('p24_ord_coral', 'Deep CORAL, order spectrum'), ('p24_ord_gdro', 'Group DRO, order spectrum')]
T12 = [[nm] + [seqcell(q[k]) for k in ('1', '5', '10', 'seq')] for nm, q in
       [('ReliSense, training priors', L28['V1']['seq']), ('ReliSense, equal priors', L28['V1_PC']['seq'])] + [(n, NET[k]) for k, n in NETS]]
V1q, PCq = L28['V1']['seq'], L28['V1_PC']['seq']
n10 = {k: V1q[k]['n'] for k in ('1', '5', '10')}
fa10 = [NET[k]['10']['fa'] for k, _ in NETS if k != 'p23_resnet_env']
acc10 = max(NET[k]['10']['acc'] for k, _ in NETS)
mr_net = [NET[k]['seq']['mean_recordings'] for k, _ in NETS]
n_h10 = 6 * 8; wrong_h10 = round(PCq['10']['fa'] * n_h10); assert abs(PCq['10']['fa'] * n_h10 - wrong_h10) < 1e-9

z = zipfile.ZipFile('ReliSense_v3_26_figures_sources.zip')
fd = np.load(io.BytesIO(z.read('ReliSense_v3_26_sources/figure_data/phase22_figdata.npz')), allow_pickle=True)
LAB = {r['bearing']: r for r in csv.DictReader(io.StringIO(z.read('ReliSense_v3_26_sources/figure_data/labels_32.csv').decode()))}
P, Y, B = fd['prob'], fd['label'], fd['bearing']; TAU = {b: fd['tau'][i, 2] for i, b in enumerate(fd['tau_bearings'])}
OK = P.argmax(1) == Y; ACC07 = P.max(1) >= np.array([TAU[b] for b in B])
BS = sorted(set(B)); NREC = {b: int((B == b).sum()) for b in BS}; PCB = L28['V1_PC']['per_bearing']
assert abs(OK.mean() - S28['PB LOBO']['V1']['acc']) < 1e-9
assert all(abs(OK[B == b].mean() - L28['V1']['per_bearing'][b]) < 1e-9 for b in BS)
def grp(sel):
    m = np.isin(B, sel); pc = sum(PCB[b] * NREC[b] for b in sel) / sum(NREC[b] for b in sel)
    return {'n': len(sel), 'acc': OK[m].mean(), 'pc': pc, 'cov': ACC07[m].mean(), 'sel': OK[m & ACC07].mean()}
H = [b for b in BS if LAB[b]['origin'] == 'healthy']; ART = [b for b in BS if LAB[b]['origin'] == 'artificial']
REAL = [b for b in BS if LAB[b]['origin'] == 'real']; DMG = ART + REAL
FAT = [b for b in REAL if LAB[b]['mode'].startswith('fatigue')]
PLS = [b for b in REAL if LAB[b]['mode'].startswith('plastic') and LAB[b]['characteristic'] == 'single point']
PLD = [b for b in REAL if LAB[b]['characteristic'] == 'distributed']; assert PLS == ['KA15'] and PLD == ['KA30']
E1 = [b for b in DMG if LAB[b]['extent'] == '1']; E23 = [b for b in DMG if LAB[b]['extent'] in ('2', '3')]
RE23 = [b for b in REAL if LAB[b]['extent'] in ('2', '3')]
G = {'H': grp(H), 'ART': grp(ART), 'FAT': grp(FAT), 'PLS': grp(PLS), 'PLD': grp(PLD), 'E1': grp(E1), 'E23': grp(E23),
     'RE23': grp(RE23), 'ALL': grp(BS), 'NO30': grp([b for b in BS if b != 'KA30'])}
def drow(name, g): return [name, str(g['n']), f1(g['acc']), f1(g['pc']), f'{f1(g["cov"])} / {f1(g["sel"])}']
T13 = [drow('Healthy', G['H']), drow('Artificial damage', G['ART']), drow('Real damage, fatigue pitting', G['FAT']),
       drow('Real damage, plastic deformation, single point (KA15)', G['PLS']),
       drow('Real damage, plastic deformation, distributed (KA30)', G['PLD']),
       drow('Damage extent level 1', G['E1']), drow('Damage extent levels 2 and 3', G['E23']),
       drow('All bearings', G['ALL']), drow('All bearings except KA30', G['NO30'])]
h30 = (B == 'KA30'); ka30_acc_n = int(ACC07[h30].sum()); ka30_ok_acc = int((ACC07 & OK)[h30].sum())
assert ka30_ok_acc * 2 < ka30_acc_n                                     # most accepted KA30 recordings wrong

def mb(S, f): return float(np.mean([S[s][f]['bacc'] for s in NAMES]))
def crow(name, S, Lk, f):
    k = Lk[f]['kappa_0.7']
    return [name, f'{100 * mb(S, f):.1f}'] + [f1(S[s][f]['bacc']) for s in ('PB L8_c0', 'PB LOBO', 'CWRU 1SIZE', 'HUST 1TYPE')] + \
           [f'{f1(k["sel_acc"])} / {f1(k["fa"])}']
L27 = R27['paderborn']['lobo']
T14 = [crow('ReliSense', S28, L28, 'V1'), crow('ReliSense, equal priors', S28, L28, 'V1_PC'),
       crow('Condition-wise standardization', S27, L27, 'CW'), crow('Multiband kinematic branch (76)', S28, L28, 'MB'),
       crow('Comb-separated kinematic branch (40)', S28, L28, 'NF'), crow('Comb-separated + slip-consistent (52)', S28, L28, 'NFS'),
       crow('Comb-separated + slip-consistent, equal priors', S28, L28, 'NFS_PC')]
assert S27['PB LOBO']['V1']['bacc'] == S28['PB LOBO']['V1']['bacc'] and R28['adoption']['adopted'] is False and R27['adoption']['adopted'] is False
cov14 = [L28[f]['kappa_0.7']['coverage'] for f in ('V1', 'V1_PC', 'MB', 'NF', 'NFS', 'NFS_PC')] + [L27['CW']['kappa_0.7']['coverage']]
d_sel = L28['V1']['kappa_0.7']['sel_acc'] - L28['V1_PC']['kappa_0.7']['sel_acc']
ct = L28['V1']['kappa_0.7_condition_thresholds']['sel_acc']
p_rf = L28['V1']['vs_RF']; assert (p_rf['better'], p_rf['worse']) == (7, 11)
def b(f, s): return f1(S28[s][f]['bacc'])

# ------------------------------------------------------------------ text
EDITS = [   # (anchor contained in the paragraph, old, new)
 ('Paper A, version', 'Paper A, version 2.6, 8 October 2026', 'Paper A, version 2.7, 8 October 2026'),
 ('Abstract—', 'with 11.2% missed faults, against 18.1% and 28.0% for the envelope residual network.',
  f'with 11.2% missed faults, and deciding from ten consecutive recordings lowered them to {f1(PCq["10"]["fa"])}% with '
  f'{f1(PCq["10"]["missed"])}% missed faults.'),
 ('1) What the Two Branches Contribute', 'less accurate on 11 (Wilcoxon p = 0.73)', f'less accurate on 11 (Wilcoxon p = {p_rf["p"]:.2f})'),
 ('4) Deployment', 'Referrals concentrated on a few bearings, and whether a repeated measurement resolves them was not evaluated.',
  'Referrals concentrated on a few bearings. Decisions from several recordings of the same operating condition are '
  'evaluated in part 5; whether a measurement at another speed or load resolves a referral was not evaluated.'),
 ('The study has limitations.', 'Features that require the agreement of several harmonics did not improve the diagnosis, and '
  'the equal-prior decision was evaluated after the main results were known.',
  'Features that require the agreement of several harmonics, condition-wise standardization, a multiband kinematic branch '
  'and comb-separated and slip-consistent fault features did not improve the diagnosis overall (Section VI-C, part 7). The '
  'equal-prior decision and the decisions from several recordings were evaluated after the main results were known, and the '
  'latter rest on six healthy bearings.'),
 ('This article addressed the diagnosis', 'fewer of both than the envelope residual network (18.1% and 28.0%).',
  f'fewer of both than the envelope residual network (18.1% and 28.0%). With ten consecutive recordings, the false alarms '
  f'fell further to {f1(PCq["10"]["fa"])}% with {f1(PCq["10"]["missed"])}% missed faults, whereas every compared network '
  f'still raised at least {f1(min(NET[k]["10"]["fa"] for k, _ in NETS))}%.'),
 ('To compare abstention at equal coverage', 'with 13.1%–23.1% false alarms (Table X).',
  'with 13.1%–23.1% false alarms (Section VI-C, part 2).'),
 ('The balance can also be shifted before abstention', 'All other results in this article use the training priors, which gave the fewest missed faults.',
  'Unless labelled otherwise, all other results in this article use the training priors, which gave the fewest missed faults.'),
 ('This article addressed the diagnosis', 'Healthy bearings that resemble damage, weak real damage and operation at low speed',
  'Healthy bearings that resemble damage, weak or distributed real damage and operation at low speed'),
]

PART5 = ('5) Decisions From Several Recordings: ',
 'A monitored bearing is measured repeatedly, so a diagnosis need not rest on a single recording. Table XII combines '
 'consecutive recordings of one bearing by summing the logarithms of their posteriors. The recordings are taken in their '
 'stored order, in which the 20 recordings of each operating condition follow each other, and the same rule was applied to '
 'the stored posteriors of every network. A sequential variant adds recordings until the combined confidence reaches 0.95, '
 f'with at most ten. With training priors, combining recordings lowered the missed faults from {f1(V1q["1"]["missed"])}% to '
 f'{f1(V1q["5"]["missed"])}% with five recordings but left the false alarms at '
 f'{f1(min(V1q[k]["fa"] for k in ("1", "5", "10")))}%–{f1(max(V1q[k]["fa"] for k in ("1", "5", "10")))}%, because most '
 'recordings of K004 and K005 were diagnosed as damaged consistently; combining recordings removes random errors, not '
 f'consistent ones. With equal priors, ten recordings, or 40 s of signal, gave {f1(PCq["10"]["acc"])}% accuracy with '
 f'{f1(PCq["10"]["fa"])}% false alarms and {f1(PCq["10"]["missed"])}% missed faults, and the sequential rule gave '
 f'{f1(PCq["seq"]["acc"])}%, {f1(PCq["seq"]["fa"])}% and {f1(PCq["seq"]["missed"])}% with '
 f'{PCq["seq"]["mean_recordings"]:.1f} recordings on average. No network gained as much. With ten recordings, the most '
 f'accurate networks reached {f1(acc10)}%, and all except the envelope ResNet raised {f1(min(fa10))}%–{f1(max(fa10))}% '
 f'false alarms; the envelope ResNet raised {f1(NET["p23_resnet_env"]["10"]["fa"])}% but missed '
 f'{f1(NET["p23_resnet_env"]["10"]["missed"])}% of the faults. These decision rules were evaluated after the main results '
 f'were known. With ten recordings, the false-alarm rate rests on {n_h10} decisions for the six healthy bearings, '
 f'{["no", "one", "two", "three", "four", "five"][wrong_h10]} of them wrong, and decisions from one bearing are not '
 'independent, so this rate is an uncertain estimate.')
T12_SPEC = ('TABLE XII', 'Decisions From Several Consecutive Recordings Under LOBO: Accuracy / False Alarms / Missed Faults (%)',
 ['Method', '1 recording', '5 recordings', '10 recordings', 'Sequential rule'], T12,
 f'Posteriors of consecutive recordings of one bearing combined by the sum of their logarithms; {n10["1"]}, {n10["5"]} and '
 f'{n10["10"]} decisions with one, five and ten recordings. Sequential rule: recordings added until the combined confidence reached 0.95, at '
 f'most 10; mean number of recordings {V1q["seq"]["mean_recordings"]:.1f} and {PCq["seq"]["mean_recordings"]:.1f} for '
 f'ReliSense with training and equal priors, {min(mr_net):.1f}–{max(mr_net):.1f} for the networks. Networks: 1 seed; '
 'WDCNN posteriors were not stored.')

PART6 = ('6) Damage Type and Extent: ',
 'Table XIII groups the LOBO results by the damage properties documented by the dataset authors [10]. The diagnosis '
 'depended more on the extent and kind of damage than on whether the damage was artificial or real. Artificial damage and '
 f'real fatigue pitting were diagnosed with similar accuracy, {f1(G["ART"]["acc"])}% and {f1(G["FAT"]["acc"])}%, so the '
 f'transfer from artificial to real damage held for the fatigue damage of {len(FAT)} of the {len(REAL)} real bearings. Damage '
 f'of extent levels 2 and 3 was diagnosed at {f1(G["E23"]["acc"])}%, against {f1(G["E1"]["acc"])}% for level 1, and the '
 f'{["", "one", "two", "three", "four"][len(RE23)]} real bearings of levels 2 and 3 at {f1(G["RE23"]["acc"])}%. The two '
 'bearings with plastic deformation by indentations were the weakest damaged bearings, KA15 at '
 f'{f1(G["PLS"]["acc"])}% and KA30 at {f1(G["PLD"]["acc"])}%. KA30 is the only single-damage bearing whose damage is '
 'documented as distributed. Distributed damage does not produce the single impact per rolling-element passage that the '
 'signal model (3) assumes, so it lies outside the scope of the fault-order features. Abstention referred '
 f'{f1(1 - G["PLD"]["cov"])}% of its recordings, but most of the accepted ones were still wrong. Without KA30, the LOBO '
 f'accuracy was {f1(G["NO30"]["acc"])}%, and abstention at κ = 0.7 gave {f1(G["NO30"]["sel"])}% selective accuracy at '
 f'{f1(G["NO30"]["cov"])}% coverage. All other results include KA30, because the damage type of a monitored bearing is not '
 'known in advance.')
T13_SPEC = ('TABLE XIII', 'LOBO Results of ReliSense by Documented Damage Properties (%)',
 ['Group', 'Bearings', 'Accuracy, training priors', 'Accuracy, equal priors', 'κ = 0.7: coverage / selective accuracy'], T13,
 'Damage properties as documented for the dataset [10]; the extent level increases with the size of the damage. '
 'Recordings pooled within each group. Equal priors: Section VI-C, part 2. κ = 0.7: training-only thresholds, training priors.')

PART7 = ('7) Changes Tested Under an Adoption Rule Fixed in Advance: ',
 'After the main results were known, four changes to the features, aimed at the errors described above, were tested on the '
 'same splits. Because several changes were tried, the adoption rule was fixed before the runs. The change with the highest '
 'mean balanced accuracy over the eleven bearing-wise settings would replace ReliSense only if it was also not less accurate '
 'under LOBO and lost at most 0.5 points of selective accuracy at κ = 0.7. Condition-wise standardization scaled each feature '
 'with the statistics of the training recordings of the same operating condition. A multiband kinematic branch read the 15 '
 'kinematic features in all four demodulation variants, 76 features in all. A comb-separated branch addressed the shaft '
 'harmonics. For the 6203 bearing, BPFO = 3 + 0.054 and BPFI = 5 − 0.054 orders, and every fault window up to the third '
 'harmonic contains an integer shaft order, whereas bearing lines are shifted from integer orders by slip [2]. This branch '
 'therefore read the fault windows without the bins within 0.03 orders of an integer order and added the strength of the '
 'integer-order line as separate features, 40 features in all. Slip-consistent comb features searched, for each fault order, '
 'one slip common to its harmonics and measured how far this comb stood out from the other integer orders (12 more features). '
 'Table XIV gives the results. No change met the rule. The highest mean balanced accuracy came from the equal-prior decision '
 f'of ReliSense itself, {100 * mb(S28, "V1_PC"):.1f}% against {100 * mb(S28, "V1"):.1f}%, but it lowered the selective '
 f'accuracy at κ = 0.7 by {100 * d_sel:.1f} points, so ReliSense was kept unchanged. The slip-consistent features helped where '
 'the test bearings differed in size or type from the training bearings. They raised the balanced accuracy on CWRU with one '
 f'training size from {b("V1", "CWRU 1SIZE")}% to {b("NFS", "CWRU 1SIZE")}% and on HUST with one training type from '
 f'{b("V1", "HUST 1TYPE")}% to {b("NFS", "HUST 1TYPE")}%, but lowered it on Paderborn L8 at condition 0 from '
 f'{b("V1", "PB L8_c0")}% to {b("NFS", "PB L8_c0")}%. The multiband branch lost '
 f'{100 * (S28["PB L8_c0"]["V1"]["bacc"] - S28["PB L8_c0"]["MB"]["bacc"]):.1f} points there. Abstention thresholds set '
 'separately for each operating condition did not help either; they lowered the selective accuracy of ReliSense at κ = 0.7 '
 f'from {f1(L28["V1"]["kappa_0.7"]["sel_acc"])}% to {f1(ct)}%.')
T14_SPEC = ('TABLE XIV', 'Changes Tested With an Adoption Rule Fixed in Advance: Balanced Accuracy (%)',
 ['Configuration', 'Mean, 11 settings', 'Paderborn L8, cond. 0', 'Paderborn LOBO', 'CWRU 1SIZE', 'HUST 1TYPE',
  'LOBO, κ = 0.7: selective accuracy / false alarms'], T14,
 'Mean over Paderborn L8 (condition 0 and all), L10, A2R and LOBO, CWRU LOSO and 1SIZE, HUST LOTO and 1TYPE, and Ottawa LOPO '
 'and 1PROF. Each change replaces the kinematic branch and keeps the envelope branch, except condition-wise standardization, '
 f'which was applied to both branches. Last column: training-only thresholds; coverage {f1(min(cov14))}%–{f1(max(cov14))}%.')
NEW_PARTS = [(PART5, T12_SPEC), (PART6, T13_SPEC), (PART7, T14_SPEC)]
FR = [[2.4, 1.9, 1.9, 1.9, 1.9], [3.6, 0.9, 1.4, 1.4, 1.9], [2.6, 1.1, 1.2, 1.2, 1.1, 1.1, 1.7]]   # relative column widths
INSERT_AFTER = 'Training priors. False alarms: healthy segments accepted'   # footnote of Table XI, end of part 4


# ------------------------------------------------------------------ docx helpers (as make_A26)
def all_pars(d):
    ps = list(d.paragraphs)
    for t in d.tables:
        for r in t.rows:
            for c in r.cells: ps += c.paragraphs
    return ps


def replace_across_runs(p, old, new):
    runs = p.runs; full = ''.join(r.text for r in runs); k = full.find(old)
    assert k >= 0 and full.count(old) == 1, (old[:70], full[:90])
    pos = 0; first = True; end = k + len(old)
    for r in runs:
        a, b_ = pos, pos + len(r.text); pos = b_
        if b_ <= k or a >= end: continue
        s, e = max(k, a) - a, min(end, b_) - a
        if first: r.text = r.text[:s] + new + r.text[e:]; first = False
        else: r.text = r.text[:s] + r.text[e:]


def find(d, anchor, exact=False):
    hits = [p for p in d.paragraphs if (p.text == anchor if exact else p.text.startswith(anchor) or anchor in p.text)]
    assert len(hits) == 1, (anchor, len(hits)); return hits[0]


def strip(el):
    for tag in ('hyperlink', 'bookmarkStart', 'bookmarkEnd', 'proofErr'):
        for x in el.findall(W + tag): el.remove(x)


def new_par_after(anchor_el, template_p, text):
    el = copy.deepcopy(template_p._p)
    rs = el.findall(W + 'r'); keep = max(rs, key=lambda r: len(''.join(t.text or '' for t in r.iter(W + 't'))))
    for r in rs:
        if r is not keep: el.remove(r)
    strip(el); anchor_el.addnext(el); p = Paragraph(el, template_p._parent); p.runs[0].text = text
    return p


def new_part_after(anchor_el, template_p, head, text):
    """paragraph with an italic run-in heading (first run of the template) and a plain body run."""
    el = copy.deepcopy(template_p._p); rs = el.findall(W + 'r')
    body = max(rs[1:], key=lambda r: len(''.join(t.text or '' for t in r.iter(W + 't'))))
    for r in rs:
        if r is not rs[0] and r is not body: el.remove(r)
    strip(el); anchor_el.addnext(el); p = Paragraph(el, template_p._parent)
    assert p.runs[0].italic and not p.runs[1].italic
    p.runs[0].text = head; p.runs[1].text = text
    return p


def setc(cell, text):
    ps = cell.paragraphs
    for extra in ps[1:]: extra._p.getparent().remove(extra._p)
    p = ps[0]; rs = p.runs
    if not rs: p.add_run(text); return
    rs[0].text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)


def build_table(tmpl_tbl, header, rows, fr=None):
    t = copy.deepcopy(tmpl_tbl._tbl); n = len(header)
    grid = t.find(W + 'tblGrid'); cols = grid.findall(W + 'gridCol')
    total = sum(int(c.get(W + 'w')) for c in cols); fr = fr or [1 / n] * n; ws = [int(total * x / sum(fr)) for x in fr]
    for c in cols[n:]: grid.remove(c)
    while len(grid.findall(W + 'gridCol')) < n: grid.append(copy.deepcopy(cols[0]))
    for c, w in zip(grid.findall(W + 'gridCol'), ws): c.set(W + 'w', str(w))
    trs = t.findall(W + 'tr'); hdr, body = trs[0], trs[1]
    for tr in trs[2:]: t.remove(tr)
    for tr in (hdr, body):
        for tc in tr.findall(W + 'tc')[n:]: tr.remove(tc)
        while len(tr.findall(W + 'tc')) < n: tr.append(copy.deepcopy(tr.findall(W + 'tc')[-1]))
        for tc, w in zip(tr.findall(W + 'tc'), ws):
            tcw = tc.find(W + 'tcPr/' + W + 'tcW')
            if tcw is not None: tcw.set(W + 'w', str(w)); tcw.set(W + 'type', 'dxa')
    t.remove(body)
    for _ in rows: t.append(copy.deepcopy(body))
    table = Table(t, tmpl_tbl._parent)
    for j, h in enumerate(header): setc(table.rows[0].cells[j], h)
    for i, row in enumerate(rows):
        for j, v in enumerate(row): setc(table.rows[i + 1].cells[j], v)
    return table


def mark(p, oldtext):
    new = p.text; A = re.findall(r'\S+\s*', oldtext); Bw = re.findall(r'\S+\s*', new)
    runs = list(p.runs); head = runs[0] if runs and runs[0].italic and oldtext.startswith(runs[0].text) and new.startswith(runs[0].text) else None
    tmpl = copy.deepcopy(max([r for r in runs if r is not head], key=lambda r: len(r.text))._r)
    if head is not None:
        A = re.findall(r'\S+\s*', oldtext[len(head.text):]); Bw = re.findall(r'\S+\s*', new[len(head.text):])
    for r in runs:
        if r is not head: r._r.getparent().remove(r._r)
    def add(t, col=None, strike=False):
        r = p.add_run(t); rpr = tmpl.find(W + 'rPr')
        if rpr is not None: r._r.insert(0, copy.deepcopy(rpr))
        if col: r.font.color.rgb = col
        if strike: r.font.strike = True
    for op, a1, a2, b1, b2 in difflib.SequenceMatcher(None, A, Bw, autojunk=False).get_opcodes():
        if op == 'equal': add(''.join(A[a1:a2]))
        if op in ('delete', 'replace'): add(''.join(A[a1:a2]), RED, True)
        if op in ('insert', 'replace'): add(''.join(Bw[b1:b2]), BLUE)


def blue(p):
    for r in p.runs: r.font.color.rgb = BLUE


def build(marked):
    d = Document(src)
    old = {p._p: p.text for p in all_pars(d)}
    for a, o, n in EDITS: replace_across_runs(find(d, a), o, n)
    cap_no = find(d, 'TABLE VI', exact=True)
    P = d.paragraphs; i6 = [k for k, p in enumerate(P) if p.text == 'TABLE VI'][0]; cap_title = P[i6 + 1]
    foot = find(d, 'Last column: selective accuracy / false alarms with training-only thresholds')
    tmpl_part = find(d, '4) Deployment'); tmpl_tbl = d.tables[5]
    pos = find(d, INSERT_AFTER)._p
    for (head, text), (no, title, hdr, rows, fn), fr in zip([p_ for p_, _ in NEW_PARTS], [t_ for _, t_ in NEW_PARTS], FR):
        q = new_part_after(pos, tmpl_part, head, text)
        q1 = new_par_after(q._p, cap_no, no); q2 = new_par_after(q1._p, cap_title, title)
        tb = build_table(tmpl_tbl, hdr, rows, fr); q2._p.addnext(tb._tbl); pos = new_par_after(tb._tbl, foot, fn)._p
    if marked:
        for p in all_pars(d):
            if p._p not in old: blue(p)
            elif p.text != old[p._p]: mark(p, old[p._p])
        P = d.paragraphs; k = [i for i, p in enumerate(P) if p.text.startswith('Revision key')]
        if k: P[k[0]]._p.getparent().remove(P[k[0]]._p)
        key = d.paragraphs[0].insert_paragraph_before('')
        key.add_run('Revision key (v2.6 → v2.7):  ')
        x = key.add_run('red strikethrough = deletion'); x.font.color.rgb = RED; x.font.strike = True
        key.add_run('   |   '); x = key.add_run('blue = addition'); x.font.color.rgb = BLUE
    return d


build(False).save(out_clean); build(True).save(out_marked); print('ok')
